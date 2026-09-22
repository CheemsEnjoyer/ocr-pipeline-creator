"""Чтение текстового слоя документа: PDF, DOCX и обычный текст."""

import io

from ..core.errors import DomainError, InvalidDocument


def digital_text(content: bytes, filename: str, mime: str):
    try:
        if mime == "application/pdf" or filename.lower().endswith(".pdf"):
            from pypdf import PdfReader
            text = "\n\n".join(page.extract_text() or "" for page in PdfReader(io.BytesIO(content)).pages)
            if not text.strip():
                raise InvalidDocument("В PDF нет текстового слоя. Выберите источник «Сканы / изображения».")
            return text.strip()
        if filename.lower().endswith(".docx"):
            from docx import Document
            doc = Document(io.BytesIO(content))
            rows = [paragraph.text for paragraph in doc.paragraphs]
            rows.extend("\t".join(cell.text for cell in row.cells) for table in doc.tables for row in table.rows)
            return "\n".join(rows).strip()
        return content.decode("utf-8-sig").strip()
    except DomainError:
        raise
    except Exception as error:
        raise InvalidDocument("Не удалось прочитать документ. Для изображений выберите источник «Сканы / изображения», для текста используйте UTF-8.") from error
