import io
import json
import os
import tempfile
import unittest
from uuid import UUID, uuid4
from contextlib import contextmanager
from pathlib import Path
from unittest.mock import patch

import httpx
from botocore.exceptions import ClientError
from fastapi.testclient import TestClient
from sqlalchemy import create_engine, inspect, select, update
from sqlalchemy import text as sql

from sqlalchemy.schema import CreateSchema, DropSchema

from backend.core.config import JobSettings, MAX_FILE_SIZE, proxy_options
from backend.db.domain.models import AdminSession, APIKey, Document, LoginGuard, ProcessingJob, TaskOutbox
from backend.db.infra.engine import configured_database_url, open_database
from backend.handler.storage import S3Storage, TemporaryFileResponse
from backend.server import create_app
from backend.service.throttle import FAILED_LOGIN_WINDOW, MAX_FAILED_LOGINS
from backend.worker.dispatcher import dispatch_once
from backend.worker.tasks import execute_job

ADMIN_LOGIN = "admin"
ADMIN_PASSWORD = "test-password-123"
ADMIN_CREDENTIALS = {"login": ADMIN_LOGIN, "password": ADMIN_PASSWORD}


class FakeS3Client:
    def __init__(self):
        self.objects = {}

    def put_object(self, Bucket, Key, Body, ContentType):
        self.objects[(Bucket, Key)] = (Body, ContentType)

    def get_object(self, Bucket, Key):
        if (Bucket, Key) not in self.objects:
            raise ClientError({"Error": {"Code": "NoSuchKey"}}, "GetObject")
        return {"Body": io.BytesIO(self.objects[(Bucket, Key)][0])}

    def delete_object(self, Bucket, Key):
        self.objects.pop((Bucket, Key), None)

    def close(self):
        pass


class APITests(unittest.TestCase):
    def setUp(self):
        self.temporary = tempfile.TemporaryDirectory()
        self.addCleanup(self.temporary.cleanup)
        self.storage = Path(self.temporary.name)
        self.database_url = self.database_for(self.storage)
        self.calls = []
        self.fail_upstream = False
        self.refuse_upstream = False
        self.environment = patch.dict("os.environ", {"AUTH_PROVIDER": "local", "LITELLM_BASE_URL": "http://litellm.test/v1", "LITELLM_API_KEY": "test-key", "OCR_ADMIN_LOGIN": ADMIN_LOGIN, "OCR_ADMIN_PASSWORD": ADMIN_PASSWORD})
        self.environment.start()
        self.addCleanup(self.environment.stop)
        self.s3_client = FakeS3Client()
        self.originals = S3Storage("test-bucket", "test/originals", self.s3_client)
        s3_factory = patch("backend.core.lifespan.S3Storage.from_env", return_value=self.originals)
        s3_factory.start()
        self.addCleanup(s3_factory.stop)
        self.app = self.make_app()
        self.client = TestClient(self.app)
        self.client.__enter__()
        self.addCleanup(self.client.__exit__, None, None, None)
        self.sign_in(self.client)

    def upstream(self, request):
        self.calls.append(request)
        if self.refuse_upstream:
            raise httpx.ConnectError("Соединение отклонено", request=request)
        if self.fail_upstream:
            return httpx.Response(503)
        if request.url.path == "/v1/models":
            return httpx.Response(200, json={"data": [{"id": "vision"}, {"id": "extract"}, {"id": "vision"}]})
        if request.url.host == "ocr.test":
            return httpx.Response(200, json={"page_content": "Распознанный текст"})
        body = json.loads(request.content)
        content = "Распознанный текст" if body["model"] == "vision" else '{"total": 1500, "verified": false}'
        if body.get("response_format", {}).get("type") == "json_schema":
            properties = body["response_format"]["json_schema"]["schema"]["properties"]
            content = json.dumps({name: "1500" if name == "total" else None for name in properties})
        return httpx.Response(200, json={"choices": [{"finish_reason": "stop", "message": {"content": content}}]})

    def make_app(self):
        return create_app(self.database_url, self.storage, httpx.MockTransport(self.upstream))

    def database_for(self, storage):
        target = os.getenv("TEST_DATABASE_URL")
        if not target:
            return f"sqlite:///{(storage / 'test.sqlite3').as_posix()}"
        # Each test owns a schema; no application or other test data is modified.
        engine, _ = open_database(target)
        schema = "ocr_test_" + uuid4().hex
        with engine.begin() as connection:
            connection.execute(CreateSchema(schema))

        def cleanup():
            try:
                with engine.begin() as connection:
                    connection.execute(DropSchema(schema, cascade=True))
            finally:
                engine.dispose()

        self.addCleanup(cleanup)
        return engine.url.update_query_dict({"options": f"-csearch_path={schema}"})

    def sign_in(self, client):
        response = client.post("/api/auth/login", json=ADMIN_CREDENTIALS)
        self.assertEqual(response.status_code, 200, response.text)
        return client

    @contextmanager
    def restarted_app(self):
        # Новый жизненный цикл приложения на той же базе, с новым входом администратора.
        with TestClient(self.make_app()) as client:
            yield self.sign_in(client)

    def upload(self, content=b'{"total": 1250}', filename="test.json", mime="application/json", pipeline=None):
        return self.client.post("/api/pipeline/run", files={"file": (filename, content, mime)}, data={"pipeline": json.dumps(pipeline or {"source": "document", "name": "Тест"})})

    def test_schema_documents_only_the_integration_api(self):
        """В /api/docs отдаётся контракт выгрузки во внешние учётные системы — и только он.

        Маршруты интерфейса продолжают работать: они скрыты из схемы, а не выключены.
        """
        schema = self.client.get("/api/openapi.json")
        self.assertEqual(schema.status_code, 200, schema.text)
        documented = {(method.upper(), path) for path, operations in schema.json()["paths"].items() for method in operations}
        self.assertEqual(documented, {
            ("GET", "/api/v1/pipelines"),
            ("POST", "/api/v1/pipeline/run"),
            ("POST", "/api/v1/pipelines/{pipeline_id}/run"),
            ("GET", "/api/v1/jobs/{job_id}"),
            ("POST", "/api/v1/jobs/{job_id}/retry"),
            ("GET", "/api/v1/documents"),
            ("GET", "/api/v1/documents/{document_id}"),
        })
        self.assertEqual(self.client.get("/api/health").status_code, 200)
        self.assertEqual(self.client.get("/api/pipelines").status_code, 200)

    def test_pipeline_catalog_and_run(self):
        self.assertEqual(self.client.get("/api/pipelines").json(), {"pipelines": []})
        config = {"id": "invoices_2026", "name": "Catalog test", "source": "scans", "ocr": {"provider": "litellm", "model": "vision", "temperature": 0.4}, "extraction": {"mode": "prompt", "model": "extract", "prompt": "JSON", "temperature": 0.7, "max_tokens": 512}}
        created = self.client.post("/api/pipelines", json=config)
        self.assertEqual(created.status_code, 201, created.text)
        pipeline = created.json()["pipeline"]
        self.assertEqual(UUID(pipeline["id"]).version, 4)
        listed = self.client.get("/api/pipelines").json()["pipelines"]
        self.assertEqual(listed, [pipeline])
        self.assertEqual(self.upload(b"image", "scan.png", "image/png", listed[0]).status_code, 200)
        with self.restarted_app() as restarted:
            self.assertEqual(restarted.get("/api/pipelines").json()["pipelines"], listed)
        changed = {**pipeline, "name": "Updated"}
        self.assertEqual(self.client.patch(f"/api/pipelines/{pipeline['id']}", json=changed).status_code, 200)
        self.assertEqual(self.client.get("/api/pipelines").json()["pipelines"][0]["name"], "Updated")
        self.assertEqual(self.client.delete(f"/api/pipelines/{pipeline['id']}").status_code, 200)
        self.assertEqual(self.client.get("/api/pipelines").json()["pipelines"], [])

    def test_generated_pipeline_uuid_and_api_access(self):
        config = {"name": "Invoices", "source": "document", "extraction": None}
        created = self.client.post("/api/pipelines", json=config)
        self.assertEqual(created.status_code, 201, created.text)
        pipeline_id = created.json()["pipeline"]["id"]
        self.assertEqual(UUID(pipeline_id).version, 4)
        self.assertEqual(str(UUID(pipeline_id)), pipeline_id)
        duplicate = self.client.post("/api/pipelines", json={**config, "id": pipeline_id})
        self.assertEqual(duplicate.status_code, 201, duplicate.text)
        self.assertNotEqual(duplicate.json()["pipeline"]["id"], pipeline_id)
        self.assertEqual(UUID(duplicate.json()["pipeline"]["id"]).version, 4)
        self.assertEqual(self.client.patch(f"/api/pipelines/{pipeline_id}", json={**config, "id": "renamed"}).status_code, 400)
        updated = self.client.patch(f"/api/pipelines/{pipeline_id}", json={**config, "name": "Updated"})
        self.assertEqual(updated.json()["pipeline"]["id"], pipeline_id)
        key = self.issue_key(**{"name": "Integration", "pipeline_ids": [pipeline_id]}).json()["token"]
        self.client.cookies.clear()
        headers = {"Authorization": f"Bearer {key}"}
        listed = self.client.get("/api/v1/pipelines", headers=headers).json()["pipelines"]
        self.assertEqual(listed[0]["id"], pipeline_id)
        result = self.client.post(f"/api/v1/pipelines/{pipeline_id}/run", headers=headers, files={"file": ("test.txt", b"Invoice 1500", "text/plain")})
        self.assertEqual(result.status_code, 200, result.text)
        document = self.client.get(f"/api/v1/documents/{result.json()['documentId']}", headers=headers).json()["document"]
        self.assertEqual(document["pipeline_id"], pipeline_id)

    def test_pipeline_import_preserves_edits_and_deletions(self):
        legacy = {"id": "pl_legacy", "name": "Legacy", "source": "document", "extraction": None}
        for _ in range(2):
            self.assertEqual(self.client.post("/api/pipelines/import", json=[legacy]).status_code, 200)
        self.assertEqual(len(self.client.get("/api/pipelines").json()["pipelines"]), 1)
        self.client.patch("/api/pipelines/pl_legacy", json={**legacy, "name": "Server edit"})
        self.client.post("/api/pipelines/import", json=[legacy])
        self.assertEqual(self.client.get("/api/pipelines").json()["pipelines"][0]["name"], "Server edit")
        self.client.delete("/api/pipelines/pl_legacy")
        self.client.post("/api/pipelines/import", json=[legacy])
        self.assertEqual(self.client.get("/api/pipelines").json()["pipelines"], [])
        self.assertEqual(self.client.patch("/api/pipelines/missing", json=legacy).status_code, 404)
        self.assertEqual(self.client.post("/api/pipelines", json={"source": "invalid"}).status_code, 422)

    def test_create_history_edit_original_and_restart(self):
        with self.app.state.sessions() as session:
            self.assertIn("documents", inspect(session.bind).get_table_names())
        response = self.upload()
        self.assertEqual(response.status_code, 200, response.text)
        document_id = response.json()["documentId"]
        self.assertEqual(self.client.get("/api/documents").json()["documents"][0]["id"], document_id)
        before = self.client.get(f"/api/documents/{document_id}").json()["document"]
        updated = self.client.patch(f"/api/documents/{document_id}", json={"fields": {"total": 1500, "lines": [1, 2]}, "revision": 0})
        self.assertEqual(updated.status_code, 200)
        source = self.client.get(f"/api/documents/{document_id}/original")
        self.assertEqual(source.content, b'{"total": 1250}')
        self.assertIn("attachment", source.headers["content-disposition"])
        conflict = self.client.patch(f"/api/documents/{document_id}", json={"fields": {}, "revision": 0})
        self.assertEqual(conflict.status_code, 409)
        # A second application lifespan reuses the same SQLite file without SQL commands.
        with self.restarted_app() as restarted:
            after = restarted.get(f"/api/documents/{document_id}").json()["document"]
            self.assertEqual(after["fields"], {"total": 1500, "lines": [1, 2]})
            self.assertEqual(after["text"], before["text"])
            self.assertEqual(after["result"], before["result"])
            self.assertEqual(after["revision"], 1)

    def test_vision_and_extraction(self):
        response = self.upload(b"image", "scan.png", "image/png", {"source": "scans", "ocr": {"provider": "litellm", "model": "vision"}, "extraction": {"mode": "fields", "model": "extract", "fields": [{"name": "total", "description": "Сумма"}]}})
        self.assertEqual(response.status_code, 200, response.text)
        self.assertEqual(response.json()["text"], "Распознанный текст")
        row = self.client.get(f"/api/documents/{response.json()['documentId']}").json()["document"]
        self.assertEqual(row["fields"], {"total": "1500"})
        self.assertTrue(all(call.headers.get("authorization") == "Bearer test-key" for call in self.calls))
        self.assertEqual(self.calls[0].url.path, "/v1/chat/completions")
        self.assertTrue(json.loads(self.calls[0].content)["messages"][0]["content"][1]["image_url"]["url"].startswith("data:image/png;base64,"))
        self.assertEqual(json.loads(self.calls[1].content)["messages"][1]["content"], "Распознанный текст")
        self.assertEqual(json.loads(self.calls[1].content)["response_format"], {
            "type": "json_schema", "json_schema": {"name": "document_fields", "strict": True, "schema": {
                "type": "object", "properties": {"total": {"type": ["string", "null"], "description": "Сумма"}},
                "required": ["total"], "additionalProperties": False,
            }},
        })

    def test_pipeline_temperatures(self):
        for temperatures in ({}, {"ocr": 0.3, "extraction": 0.8}, {"ocr": 0, "extraction": 2}):
            with self.subTest(temperatures=temperatures):
                self.calls.clear()
                pipeline = {"source": "scans", "ocr": {"provider": "litellm", "model": "vision"}, "extraction": {"mode": "prompt", "model": "extract", "prompt": "JSON"}}
                for stage, value in temperatures.items():
                    pipeline[stage]["temperature"] = value
                response = self.upload(b"image", "scan.png", "image/png", pipeline)
                self.assertEqual(response.status_code, 200, response.text)
                self.assertEqual([json.loads(call.content)["temperature"] for call in self.calls], [temperatures.get("ocr", 0), temperatures.get("extraction", 0)])

    def test_pipeline_priority_reaches_celery_and_survives_retry(self):
        from backend.worker.tasks import enqueue_job, process_document

        pipeline = {"source": "document", "name": "Priority", "priority": 0,
                    "extraction": {"mode": "prompt", "model": "extract", "prompt": "Extract"}}
        job = self.submit_background(pipeline=pipeline)
        with self.app.state.sessions() as session:
            row = session.get(ProcessingJob, job["taskId"])
            event = session.scalar(select(TaskOutbox).where(TaskOutbox.job_id == row.id))
            self.assertEqual((row.priority, event.priority, row.pipeline["priority"]), (0, 0, 0))
        published = []
        self.assertEqual(dispatch_once(self.app.state.sessions, publisher=lambda *args: published.append(args)), 1)
        self.assertEqual(published[0][3], 0)
        with patch.object(process_document, "apply_async") as apply_async:
            enqueue_job(*published[0])
        self.assertEqual(apply_async.call_args.kwargs["priority"], 0)

        self.fail_upstream = True
        self.execute_background(job["taskId"])
        self.fail_upstream = False
        self.client.post(job["statusUrl"] + "/retry")
        with self.app.state.sessions() as session:
            retry_event = session.scalar(select(TaskOutbox).where(TaskOutbox.job_id == job["taskId"], TaskOutbox.published_at.is_(None)))
            self.assertEqual(retry_event.priority, 0)

        for invalid in (-1, 10, 1.5, "9"):
            with self.subTest(priority=invalid):
                response = self.client.post("/api/pipelines", json={"id": f"priority_{str(invalid).replace('.', '_').replace('-', 'n')}", "name": "Invalid", "source": "document", "priority": invalid})
                self.assertEqual(response.status_code, 422)

    def test_pipeline_execution_modes_and_async_concurrency(self):
        async_config = {"id": "async_limited", "name": "Async", "source": "document",
                        "allow_sync": False, "allow_async": True, "async_concurrency": 1}
        created = self.client.post("/api/pipelines", json=async_config)
        self.assertEqual(created.status_code, 201, created.text)
        async_id = created.json()["pipeline"]["id"]
        first = self.client.post("/api/pipeline/run?background=true", files={"file": ("1.txt", b"one", "text/plain")}, data={"pipeline_id": async_id})
        second = self.client.post("/api/pipeline/run?background=true", files={"file": ("2.txt", b"two", "text/plain")}, data={"pipeline_id": async_id})
        self.assertEqual((first.status_code, second.status_code), (202, 202))
        self.assertEqual(self.client.post(f"/api/pipelines/{async_id}/run?background=false", files={"file": ("x.txt", b"x", "text/plain")}).status_code, 400)

        clock = 1000
        claimed = self.app.state.processing.jobs.claim(first.json()["taskId"], clock=clock)
        self.assertIsNotNone(claimed)
        self.assertIsNone(self.app.state.processing.jobs.claim(second.json()["taskId"], clock=clock))
        with self.app.state.sessions.begin() as session:
            delayed = session.get(ProcessingJob, second.json()["taskId"])
            self.assertEqual((delayed.status, delayed.generation, delayed.next_run_at), ("queued", 1, clock + 1))
            session.get(ProcessingJob, first.json()["taskId"]).status = "succeeded"
        self.assertIsNotNone(self.app.state.processing.jobs.claim(second.json()["taskId"], generation=1, clock=clock + 2))

        sync_config = {"id": "sync_only", "name": "Sync", "source": "document", "allow_sync": True, "allow_async": False}
        created = self.client.post("/api/pipelines", json=sync_config)
        self.assertEqual(created.status_code, 201, created.text)
        sync_id = created.json()["pipeline"]["id"]
        self.assertEqual(self.client.post(f"/api/pipelines/{sync_id}/run", files={"file": ("x.txt", b"x", "text/plain")}).status_code, 200)
        self.assertEqual(self.client.post(f"/api/pipelines/{sync_id}/run?background=true", files={"file": ("x.txt", b"x", "text/plain")}).status_code, 400)

        maximum = self.client.post("/api/pipelines", json={"id": "max_concurrency", "name": "Maximum", "source": "document", "async_concurrency": 30})
        self.assertEqual(maximum.status_code, 201, maximum.text)

        for invalid in ({"allow_sync": False, "allow_async": False}, {"async_concurrency": 0}, {"async_concurrency": 31}):
            response = self.client.post("/api/pipelines", json={"id": "invalid_modes", "name": "Invalid", "source": "document", **invalid})
            self.assertEqual(response.status_code, 422, response.text)

    def test_distributed_sync_capacity_limits_global_and_client_load(self):
        from backend.core.errors import LimitExceeded
        from backend.db.domain.models import IntegrationClient
        from backend.db.infra.repositories import SyncCapacityRepository

        first_client, second_client = str(uuid4()), str(uuid4())
        with self.app.state.sessions.begin() as session:
            session.add_all([
                IntegrationClient(id=first_client, name="First", pipeline_ids=[], created_at="2026"),
                IntegrationClient(id=second_client, name="Second", pipeline_ids=[], created_at="2026"),
            ])
        settings = JobSettings(timeout=60, sync_global_limit=3, sync_client_limit=2)
        capacity = SyncCapacityRepository(self.app.state.sessions, settings)
        first = capacity.acquire(first_client, clock=1000)
        second = capacity.acquire(first_client, clock=1000)
        with self.assertRaises(LimitExceeded) as client_error:
            capacity.acquire(first_client, clock=1000)
        self.assertEqual((client_error.exception.status, client_error.exception.headers["Retry-After"]), (429, "5"))
        third = capacity.acquire(second_client, clock=1000)
        with self.assertRaises(LimitExceeded) as global_error:
            capacity.acquire(second_client, clock=1000)
        self.assertEqual(global_error.exception.status, 429)
        capacity.release(first)
        replacement = capacity.acquire(second_client, clock=1001)
        # Leases left behind by a crashed API process stop counting after timeout + grace period.
        recovered = capacity.acquire(first_client, clock=1121)
        self.assertIsInstance(recovered, str)
        for slot in (second, third, replacement, recovered):
            capacity.release(slot)

        live_capacity = self.app.state.processing.sync_capacity
        occupied = [live_capacity.acquire() for _ in range(10)]
        try:
            blocked = self.upload(pipeline={"source": "document"})
            self.assertEqual((blocked.status_code, blocked.headers.get("Retry-After")), (429, "5"))
            self.assertEqual(self.calls, [])
        finally:
            for slot in occupied:
                live_capacity.release(slot)

    def test_prompt_result_stays_text_even_if_model_returns_json(self):
        pipeline = {"source": "document", "extraction": {"mode": "prompt", "model": "extract", "prompt": "Кратко опиши документ"}}
        response = self.upload(pipeline=pipeline)
        self.assertEqual(response.status_code, 200, response.text)
        payload = json.loads(self.calls[-1].content)
        self.assertNotIn("response_format", payload)
        self.assertNotIn("только валидным JSON", payload["messages"][0]["content"])
        self.assertIn("Кратко опиши документ", payload["messages"][0]["content"])
        with self.restarted_app() as client:
            row = client.get(f"/api/documents/{response.json()['documentId']}").json()["document"]
            self.assertEqual(row["fields"], {})
            self.assertEqual(row["result"], response.json()["result"])

    def test_prompt_result_preserves_paragraphs(self):
        result = "Общее описание документа.\n\nВторой абзац с выводами."
        pipeline = {"source": "document", "extraction": {"mode": "prompt", "model": "extract", "prompt": "Опиши документ"}}
        with patch("backend.handler.litellm.complete", return_value=result) as completion:
            response = self.upload(pipeline=pipeline)
        self.assertEqual(response.status_code, 200, response.text)
        self.assertNotIn("response_format", completion.call_args.args[1])
        row = self.client.get(f"/api/documents/{response.json()['documentId']}").json()["document"]
        self.assertEqual(row["result"], result)
        self.assertEqual(row["fields"], {})

    def test_invalid_temperatures_do_not_call_upstream(self):
        for stage in ("ocr", "extraction"):
            for value in (-0.1, 2.1, "invalid"):
                with self.subTest(stage=stage, value=value):
                    pipeline = {"source": "scans", "ocr": {"provider": "litellm", "model": "vision"}, "extraction": {"mode": "prompt", "model": "extract"}}
                    pipeline[stage]["temperature"] = value
                    response = self.upload(b"image", "scan.png", "image/png", pipeline)
                    self.assertEqual(response.status_code, 400)
        self.assertEqual(self.calls, [])

    def test_ocr_service_and_failure(self):
        options = {"model_name": "deepseek-ai/DeepSeek-OCR", "force_ocr": "True"}
        pipeline = {"source": "scans", "ocr": {"provider": "service", "url": "http://ocr.test/recognize", "options": options}}
        response = self.upload(b"image", "scan.png", "image/png", pipeline)
        self.assertEqual(response.status_code, 200)
        self.assertEqual(response.json()["result"], "Распознанный текст")
        # Поля сервиса уходят в том же multipart-запросе, что и файл.
        body = self.calls[0].content.decode("utf-8", "replace")
        self.assertIn('name="model_name"', body)
        self.assertIn("deepseek-ai/DeepSeek-OCR", body)
        self.assertIn('name="force_ocr"', body)
        self.assertIn('name="file"; filename="scan.png"', body)
        self.fail_upstream = True
        failed = self.upload(b"image", "scan.png", "image/png", pipeline)
        self.assertEqual(failed.status_code, 502)
        self.assertEqual(len(self.client.get("/api/documents").json()["documents"]), 1)
        self.assertEqual(len(self.s3_client.objects), 1)

    def test_health_is_public_and_hides_configuration(self):
        # health открыт без входа, поэтому адреса LiteLLM и прокси в нём не раскрываются.
        health = TestClient(self.app).get("/api/health")
        self.assertEqual(health.status_code, 200)
        self.assertEqual(health.json(), {"status": "ok", "backend": "fastapi-sqlalchemy"})
        blank = dict.fromkeys(("PROXY_URL", "HTTP_PROXY", "HTTPS_PROXY", "http_proxy", "https_proxy"), "")
        with patch.dict("os.environ", {**blank, "PROXY_URL": "proxy.company.local:8080"}):
            app = self.make_app()
            with TestClient(app):
                self.assertEqual(app.state.proxy, "http://proxy.company.local:8080")

    def create_pipeline(self, name):
        response = self.client.post("/api/pipelines", json={"name": name, "source": "document", "extraction": None})
        self.assertEqual(response.status_code, 201, response.text)
        return response.json()["pipeline"]["id"]

    def issue_key(self, name, pipeline_ids, client=None):
        admin = client or self.client
        created = admin.post("/api/clients", json={"name": name, "pipeline_ids": pipeline_ids})
        if created.status_code != 201:
            return created
        return admin.post("/api/keys", json={"name": name, "client_id": created.json()["client"]["id"], "permissions": ["run", "results", "history"]})

    @staticmethod
    def send_document(client, url, data=None):
        return client.post(url, files={"file": ("doc.txt", "Текст документа".encode(), "text/plain")}, data=data or {})

    def test_requests_without_credentials_are_rejected(self):
        anonymous = TestClient(self.app)
        for method, url in (("get", "/api/pipelines"), ("get", "/api/documents"), ("get", "/api/keys"), ("get", "/api/auth/session"), ("get", "/api/litellm/models")):
            with self.subTest(url=url):
                self.assertEqual(getattr(anonymous, method)(url).status_code, 401)
        self.assertEqual(self.send_document(anonymous, "/api/pipeline/run", {"pipeline": json.dumps({"source": "document"})}).status_code, 401)
        # Пароль администратора в заголовке не работает — там принимаются только ключи интеграций.
        for header in ("Bearer wrong-key", "Bearer " + ADMIN_PASSWORD, "Basic dXNlcjpwYXNz", "Bearer " + "x" * 600):
            with self.subTest(header=header[:20]):
                self.assertEqual(anonymous.get("/api/pipelines", headers={"Authorization": header}).status_code, 401)

    def test_admin_login_session_expiry_and_logout(self):
        browser = TestClient(self.app)
        for credentials in ({"login": ADMIN_LOGIN, "password": "wrong-password"}, {"login": "root", "password": ADMIN_PASSWORD}):
            with self.subTest(login=credentials["login"]):
                wrong = browser.post("/api/auth/login", json=credentials)
                # Одинаковый ответ: по нему не понять, что именно неверно.
                self.assertEqual((wrong.status_code, wrong.json()["error"]), (401, "Неверный логин или пароль"))
        login = browser.post("/api/auth/login", json={"login": f"  {ADMIN_LOGIN} ", "password": ADMIN_PASSWORD})
        self.assertEqual(login.status_code, 200, login.text)
        cookie = login.headers["set-cookie"].lower()
        self.assertIn("httponly", cookie)
        self.assertIn("samesite=strict", cookie)
        token = browser.cookies.get("ocr_admin_session")
        with self.app.state.sessions() as session:
            rows = session.scalars(select(AdminSession)).all()
            # В базе нет ни токена сессии, ни пароля в открытом виде.
            self.assertNotIn(token, [row.token_hash for row in rows])
            self.assertFalse(any(ADMIN_PASSWORD in row.admin_hash for row in rows))
        identity = browser.get("/api/auth/session").json()
        self.assertEqual(identity["role"], "admin")
        self.assertEqual(identity["login"], ADMIN_LOGIN)
        self.assertEqual(browser.get("/api/documents").status_code, 200)
        with self.app.state.sessions.begin() as session:
            session.execute(update(AdminSession).values(expires_at=0))
        self.assertEqual(browser.get("/api/auth/session").status_code, 401)

        self.sign_in(browser)
        token = browser.cookies.get("ocr_admin_session")
        self.assertEqual(browser.post("/api/auth/logout").status_code, 200)
        # Старый cookie не оживает, даже если клиент пришлёт его снова.
        self.assertEqual(TestClient(self.app, cookies={"ocr_admin_session": token}).get("/api/auth/session").status_code, 401)

    def test_admin_password_is_generated_once_when_not_configured(self):
        storage = self.storage / "fresh"
        database = self.database_for(storage)
        with patch.dict("os.environ", {"OCR_ADMIN_LOGIN": "", "OCR_ADMIN_PASSWORD": ""}):
            with TestClient(create_app(database, storage, httpx.MockTransport(self.upstream))) as first:
                password = (storage / "admin-password.txt").read_text(encoding="utf-8").strip()
                self.assertGreaterEqual(len(password), 16)
                self.assertEqual(first.post("/api/auth/login", json={"login": "admin", "password": password}).status_code, 200)
            with TestClient(create_app(database, storage, httpx.MockTransport(self.upstream))) as second:
                self.assertEqual(second.post("/api/auth/login", json={"login": "admin", "password": password}).status_code, 200)
            self.assertEqual((storage / "admin-password.txt").read_text(encoding="utf-8").strip(), password)
        with patch.dict("os.environ", {"OCR_ADMIN_PASSWORD": "short"}), self.assertRaisesRegex(RuntimeError, "OCR_ADMIN_PASSWORD"):
            with TestClient(self.make_app()):
                pass

    def test_changing_admin_password_ends_sessions(self):
        token = self.client.cookies.get("ocr_admin_session")
        with TestClient(self.make_app(), cookies={"ocr_admin_session": token}) as same_password:
            self.assertEqual(same_password.get("/api/auth/session").status_code, 200)
        with patch.dict("os.environ", {"OCR_ADMIN_PASSWORD": "another-password-456"}), TestClient(self.make_app(), cookies={"ocr_admin_session": token}) as changed:
            self.assertEqual(changed.get("/api/auth/session").status_code, 401)
            self.assertEqual(changed.post("/api/auth/login", json=ADMIN_CREDENTIALS).status_code, 401)
            self.assertEqual(changed.post("/api/auth/login", json={"login": ADMIN_LOGIN, "password": "another-password-456"}).status_code, 200)

    def test_repeated_failed_logins_are_throttled(self):
        browser = TestClient(self.app)
        wrong = {"login": ADMIN_LOGIN, "password": "wrong-password"}
        for _ in range(MAX_FAILED_LOGINS):
            self.assertEqual(browser.post("/api/auth/login", json=wrong).status_code, 401)
        locked = browser.post("/api/auth/login", json=ADMIN_CREDENTIALS)
        self.assertEqual(locked.status_code, 429)
        self.assertIn("Повторите через 10 мин", locked.json()["error"])
        # Окно прошло — верный пароль снова работает, а счётчик сбрасывается.
        with self.app.state.sessions.begin() as session:
            guard = session.get(LoginGuard, "admin")
            guard.failures = [moment - FAILED_LOGIN_WINDOW for moment in guard.failures]
        self.assertEqual(browser.post("/api/auth/login", json=ADMIN_CREDENTIALS).status_code, 200)
        with self.app.state.sessions() as session:
            self.assertEqual(session.get(LoginGuard, "admin").failures, [])

    def test_integration_key_is_limited_to_assigned_pipelines(self):
        allowed = self.create_pipeline("Счета")
        hidden = self.create_pipeline("Договоры")
        created = self.issue_key(**{"name": "1С", "pipeline_ids": [allowed]})
        self.assertEqual(created.status_code, 201, created.text)
        token = created.json()["token"]
        self.assertTrue(token.startswith("ocr_"))
        listed = self.client.get("/api/keys").json()["keys"]
        self.assertEqual([(key["name"], key["pipeline_ids"], key["prefix"]) for key in listed], [("1С", [allowed], token[:12])])
        # Ключ показывается один раз: ни список, ни база его не содержат.
        self.assertNotIn(token, json.dumps(listed))
        with self.app.state.sessions() as session:
            self.assertNotIn(token, session.scalars(select(APIKey.token_hash)).all())

        integration = TestClient(self.app, headers={"Authorization": f"Bearer {token}"})
        self.assertEqual([pipeline["id"] for pipeline in integration.get("/api/pipelines").json()["pipelines"]], [allowed])
        by_path = self.send_document(integration, f"/api/pipelines/{allowed}/run")
        self.assertEqual(by_path.status_code, 200, by_path.text)
        self.assertEqual(by_path.json()["text"], "Текст документа")
        self.assertEqual(self.send_document(integration, "/api/pipeline/run", {"pipeline_id": allowed}).status_code, 200)
        self.assertEqual(self.send_document(integration, f"/api/pipelines/{hidden}/run").status_code, 403)
        self.assertEqual(self.send_document(integration, "/api/pipeline/run", {"pipeline_id": hidden}).status_code, 403)
        self.assertEqual(self.send_document(integration, "/api/pipeline/run", {"pipeline": json.dumps({"source": "document"})}).status_code, 403)
        for method, url in (("get", "/api/documents"), ("get", "/api/keys"), ("post", "/api/keys"), ("get", "/api/auth/session"), ("post", "/api/pipelines"), ("delete", f"/api/pipelines/{allowed}"), ("get", "/api/litellm/models")):
            with self.subTest(method=method, url=url):
                self.assertEqual(getattr(integration, method)(url).status_code, 403)
        self.assertEqual(TestClient(self.app).post("/api/auth/login", json={"login": ADMIN_LOGIN, "password": token}).status_code, 401)

    def test_versioned_api_matches_unversioned_paths(self):
        allowed = self.create_pipeline("Счета")
        hidden = self.create_pipeline("Договоры")
        token = self.issue_key(**{"name": "ERP", "pipeline_ids": [allowed]}).json()["token"]
        integration = TestClient(self.app, headers={"Authorization": "Bearer " + token})
        self.assertEqual([pipeline["id"] for pipeline in integration.get("/api/v1/pipelines").json()["pipelines"]], [allowed])
        by_path = self.send_document(integration, f"/api/v1/pipelines/{allowed}/run")
        self.assertEqual(by_path.status_code, 200, by_path.text)
        self.assertEqual(by_path.json()["text"], "Текст документа")
        self.assertEqual(self.send_document(integration, "/api/v1/pipeline/run", {"pipeline_id": allowed}).status_code, 200)
        # Ограничения ключа те же, что и на путях без версии.
        self.assertEqual(self.send_document(integration, f"/api/v1/pipelines/{hidden}/run").status_code, 403)
        self.assertEqual(self.send_document(integration, "/api/v1/pipeline/run", {"pipeline_id": hidden}).status_code, 403)
        self.assertEqual(self.send_document(integration, "/api/v1/pipeline/run", {"pipeline": json.dumps({"source": "document"})}).status_code, 403)
        self.assertEqual(TestClient(self.app).get("/api/v1/pipelines").status_code, 401)
        # Versioned API exposes only public contracts even to an administrator.
        self.assertEqual({p["id"] for p in self.client.get("/api/v1/pipelines").json()["pipelines"]}, {allowed, hidden})
        self.assertNotIn("source", self.client.get("/api/v1/pipelines").json()["pipelines"][0])
        self.assertEqual(self.client.get("/api/v1/keys").status_code, 404)

    def test_integration_key_sees_only_documents_it_processed(self):
        pipeline = self.create_pipeline("Счета")
        first = self.issue_key(**{"name": "1С", "pipeline_ids": [pipeline]}).json()
        second = self.issue_key(**{"name": "CRM", "pipeline_ids": [pipeline]}).json()
        integration = TestClient(self.app, headers={"Authorization": "Bearer " + first["token"]})
        other = TestClient(self.app, headers={"Authorization": "Bearer " + second["token"]})

        own = [
            self.send_document(integration, f"/api/v1/pipelines/{pipeline}/run").json()["documentId"],
            self.send_document(integration, "/api/pipeline/run", {"pipeline_id": pipeline}).json()["documentId"],
        ]
        foreign = self.send_document(other, f"/api/v1/pipelines/{pipeline}/run").json()["documentId"]
        manual = self.upload().json()["documentId"]

        listed = integration.get("/api/v1/documents")
        self.assertEqual(listed.status_code, 200, listed.text)
        self.assertEqual(sorted(item["id"] for item in listed.json()["documents"]), sorted(own))
        self.assertFalse(listed.json()["hasMore"])
        self.assertEqual({item["pipeline_id"] for item in listed.json()["documents"]}, {pipeline})

        detail = integration.get(f"/api/v1/documents/{own[0]}")
        self.assertEqual(detail.status_code, 200, detail.text)
        document = detail.json()["document"]
        self.assertEqual((document["text"], document["pipeline_id"], document["pipeline_name"]), ("Текст документа", pipeline, "Счета"))
        # Чужой документ неотличим от несуществующего.
        for hidden in (foreign, manual, "missing"):
            with self.subTest(document=hidden):
                self.assertEqual(integration.get(f"/api/v1/documents/{hidden}").status_code, 404)
        self.assertEqual([item["id"] for item in other.get("/api/v1/documents").json()["documents"]], [foreign])

        # Поля, исправленные оператором в истории, интеграция получает уже исправленными.
        self.assertEqual(self.client.patch(f"/api/documents/{own[0]}", json={"fields": {"total": "1 500"}, "revision": 0}).status_code, 200)
        self.assertEqual(integration.get(f"/api/v1/documents/{own[0]}").json()["document"]["fields"], {"total": "1 500"})

        # Общая история, исходный файл и правка полей ключу недоступны.
        self.assertEqual(integration.get("/api/documents").status_code, 403)
        self.assertEqual(integration.get(f"/api/documents/{own[0]}").status_code, 403)
        self.assertEqual(integration.get(f"/api/documents/{own[0]}/original").status_code, 403)
        self.assertEqual(integration.patch(f"/api/documents/{own[0]}", json={"fields": {}, "revision": 1}).status_code, 403)
        self.assertEqual(integration.get(f"/api/v1/documents/{own[0]}/original").status_code, 404)

        # Администратор по-прежнему видит всю историю, в том числе через /api/v1.
        self.assertEqual({item["id"] for item in self.client.get("/api/v1/documents").json()["documents"]}, {*own, foreign, manual})
        self.assertIsNone(self.client.get(f"/api/documents/{manual}").json()["document"]["pipeline_id"])

        # Отозванный ключ теряет доступ, а сменный ключ того же клиента сохраняет историю.
        self.client.delete(f"/api/keys/{first['key']['id']}")
        self.assertEqual(integration.get("/api/v1/documents").status_code, 401)
        replacement = self.client.post("/api/keys", json={"name": "1С — новый", "client_id": first["key"]["client_id"], "permissions": ["run", "results", "history"]}).json()["token"]
        self.assertEqual({row["id"] for row in TestClient(self.app, headers={"Authorization": "Bearer " + replacement}).get("/api/v1/documents").json()["documents"]}, set(own))

    def test_existing_documents_table_is_upgraded(self):
        storage = self.storage / "legacy"
        storage.mkdir()
        database = self.database_for(storage)
        engine = create_engine(database)
        with engine.begin() as connection:
            # Схема таблицы документов до появления pipeline_id и api_key_id.
            connection.execute(sql("CREATE TABLE documents (id VARCHAR(36) PRIMARY KEY, filename TEXT, mime_type TEXT, size INTEGER, pipeline_name TEXT, original_key TEXT, text TEXT, result TEXT, fields JSON, created_at VARCHAR(32), updated_at VARCHAR(32), revision INTEGER)"))
            connection.execute(sql("INSERT INTO documents VALUES ('old-doc', 'old.txt', 'text/plain', 3, 'Старый', 'old-doc', 'abc', 'abc', '{}', '2025-01-01T00:00:00.000Z', '2025-01-01T00:00:00.000Z', 0)"))
        engine.dispose()

        app = create_app(database, storage, httpx.MockTransport(self.upstream))
        with TestClient(app) as client:
            self.sign_in(client)
            with app.state.sessions() as session:
                inspector = inspect(session.get_bind())
                self.assertLessEqual({"pipeline_id", "api_key_id"}, {column["name"] for column in inspector.get_columns("documents")})
                indexes = {index["name"] for index in inspector.get_indexes("documents")}
                self.assertIn("idx_documents_pipeline", indexes)
                # По api_key_id не фильтруют: индекс снят, колонка осталась как аудит.
                self.assertNotIn("idx_documents_api_key", indexes)
            old = client.get("/api/documents/old-doc").json()["document"]
            self.assertEqual((old["filename"], old["fields"], old["pipeline_id"]), ("old.txt", {}, None))
            pipeline = client.post("/api/pipelines", json={"id": "new_pipeline", "name": "Новый", "source": "document", "extraction": None}).json()["pipeline"]["id"]
            token = self.issue_key(client=client, **{"name": "ERP", "pipeline_ids": [pipeline]}).json()["token"]
            integration = TestClient(app, headers={"Authorization": "Bearer " + token})
            # Старые документы остаются только у администратора.
            self.assertEqual(integration.get("/api/v1/documents").json()["documents"], [])
            self.assertEqual(integration.get("/api/v1/documents/old-doc").status_code, 404)
        # Повторный запуск на уже обновлённой базе проходит без ошибок.
        with TestClient(create_app(database, storage, httpx.MockTransport(self.upstream))) as again:
            self.assertEqual(self.sign_in(again).get("/api/documents").json()["documents"][0]["id"], "old-doc")

    def test_existing_keys_documents_and_jobs_receive_client_owner(self):
        from alembic import command
        from alembic.config import Config
        from backend.db.domain.models import IntegrationClient

        migration_storage = self.storage / "owned-migration"
        migration_storage.mkdir()
        database = self.database_for(migration_storage)
        engine = create_engine(database)
        config = Config()
        config.set_main_option("script_location", str(Path(__file__).parents[1] / "db" / "migration"))
        with engine.begin() as connection:
            config.attributes["connection"] = connection
            command.upgrade(config, "0007_result_schema")
            connection.execute(sql("INSERT INTO pipelines (id, config, created_at, updated_at, deleted) VALUES ('invoice', '{\"name\":\"Invoice\",\"source\":\"document\"}', '2025', '2025', 0)"))
            connection.execute(sql("INSERT INTO api_keys (id, name, token_hash, prefix, pipeline_ids, created_at, revoked_at) VALUES ('old-key', 'Legacy', 'hash', 'ocr_old', '[\"invoice\"]', '2025', NULL)"))
            connection.execute(sql("INSERT INTO documents (id, filename, mime_type, size, pipeline_name, original_key, text, result, fields, created_at, updated_at, revision, pipeline_id, api_key_id, result_schema) VALUES ('old-doc', 'old.txt', 'text/plain', 3, 'Invoice', 'old', 'abc', 'abc', '{}', '2025', '2025', 0, 'invoice', 'old-key', '{\"type\":\"string\"}')"))
            connection.execute(sql("INSERT INTO processing_jobs (id, status, filename, mime_type, size, original_key, pipeline, pipeline_id, api_key_id, error, created_at, updated_at, generation, attempts, next_run_at) VALUES ('old-job', 'queued', 'old.txt', 'text/plain', 3, 'old-job', '{\"name\":\"Invoice\",\"source\":\"document\"}', 'invoice', 'old-key', NULL, '2025', '2025', 0, 0, 0)"))
        with engine.begin() as connection:
            config.attributes["connection"] = connection
            command.upgrade(config, "head")
        migrated_engine, sessions = open_database(database)
        with sessions() as session:
            key = session.get(APIKey, "old-key")
            self.assertEqual((key.client_id, key.permissions), ("old-key", ["run", "results", "history"]))
            self.assertEqual((session.get(IntegrationClient, "old-key").name, session.get(IntegrationClient, "old-key").pipeline_ids), ("Legacy", ["invoice"]))
            self.assertEqual(session.get(Document, "old-doc").client_id, "old-key")
            self.assertEqual(session.get(ProcessingJob, "old-job").client_id, "old-key")
        migrated_engine.dispose()
        engine.dispose()

    def test_dangling_pipeline_reference_is_restored_as_a_deleted_pipeline(self):
        from alembic import command
        from alembic.config import Config
        from backend.db.domain.models import SavedPipeline

        migration_storage = self.storage / "dangling-migration"
        migration_storage.mkdir()
        database = self.database_for(migration_storage)
        engine = create_engine(database)
        config = Config()
        config.set_main_option("script_location", str(Path(__file__).parents[1] / "db" / "migration"))
        with engine.begin() as connection:
            config.attributes["connection"] = connection
            command.upgrade(config, "0011_job_stages")
            # Пайплайн удалён физически прежней версией: ссылка осталась висячей.
            connection.execute(sql("INSERT INTO documents (id, filename, mime_type, size, pipeline_name, original_key, text, result, fields, created_at, updated_at, revision, pipeline_id, api_key_id, result_schema) VALUES ('orphan-doc', 'old.txt', 'text/plain', 3, 'Исчезнувший', 'orphan', 'abc', 'abc', '{}', '2025', '2025', 0, 'vanished', NULL, '{\"type\":\"string\"}')"))
        with engine.begin() as connection:
            config.attributes["connection"] = connection
            command.upgrade(config, "head")
        migrated_engine, sessions = open_database(database)
        with sessions() as session:
            restored = session.get(SavedPipeline, "vanished")
            self.assertEqual((restored.deleted, restored.config["name"]), (1, "Исчезнувший"))
            # Ссылка сохранена, а значит история по пайплайну продолжает фильтроваться.
            self.assertEqual(session.get(Document, "orphan-doc").pipeline_id, "vanished")
        migrated_engine.dispose()
        engine.dispose()

    def test_key_access_can_be_changed_and_revoked(self):
        first = self.create_pipeline("Первый")
        second = self.create_pipeline("Второй")
        created = self.issue_key(**{"name": "CRM", "pipeline_ids": [first]}).json()
        key_id = created["key"]["id"]
        integration = TestClient(self.app, headers={"Authorization": "Bearer " + created["token"]})

        client_id = created["key"]["client_id"]
        updated = self.client.patch(f"/api/clients/{client_id}", json={"name": "CRM 2", "pipeline_ids": [second, second]})
        self.assertEqual(updated.status_code, 200, updated.text)
        self.assertEqual((updated.json()["client"]["name"], updated.json()["client"]["pipeline_ids"]), ("CRM 2", [second]))
        self.assertEqual(self.send_document(integration, f"/api/pipelines/{first}/run").status_code, 403)
        self.assertEqual(self.send_document(integration, f"/api/pipelines/{second}/run").status_code, 200)

        self.client.delete(f"/api/pipelines/{first}")
        for ids in ([first], ["pl_missing"]):
            with self.subTest(ids=ids):
                self.assertEqual(self.client.patch(f"/api/clients/{client_id}", json={"name": "CRM", "pipeline_ids": ids}).status_code, 422)
                self.assertEqual(self.issue_key(**{"name": "Новый", "pipeline_ids": ids}).status_code, 422)
        self.assertEqual(self.issue_key(**{"name": "   ", "pipeline_ids": [second]}).status_code, 422)

        # Удалённый пайплайн пропадает у ключа, не ломая его.
        self.client.delete(f"/api/pipelines/{second}")
        self.assertEqual(integration.get("/api/pipelines").json()["pipelines"], [])
        self.assertEqual(self.send_document(integration, f"/api/pipelines/{second}/run").status_code, 404)

        self.assertEqual(self.client.delete(f"/api/keys/{key_id}").status_code, 200)
        self.assertEqual(integration.get("/api/pipelines").status_code, 401)
        self.assertEqual(self.client.delete(f"/api/keys/{key_id}").status_code, 200)
        self.assertEqual(self.client.patch(f"/api/keys/{key_id}", json={"name": "CRM", "client_id": client_id, "permissions": ["run"]}).status_code, 404)
        self.assertEqual(self.client.delete("/api/keys/missing").status_code, 404)
        self.assertIsNotNone(self.client.get("/api/keys").json()["keys"][0]["revoked_at"])

    def test_models_and_chat(self):
        self.assertEqual(self.client.get("/api/litellm/models").json()["models"], ["extract", "vision"])
        chat = self.client.post("/api/litellm/chat", json={"model": "extract", "prompt": "Извлеки сумму", "documentText": "1500"})
        self.assertEqual(chat.status_code, 200)
        self.assertIn("1500", chat.json()["result"])

    def test_validation_and_missing_documents(self):
        self.assertEqual(self.client.get("/api/documents/missing").status_code, 404)
        self.assertEqual(self.client.get("/api/documents/missing/original").status_code, 404)
        self.assertEqual(self.client.patch("/api/documents/missing", json={"fields": [], "revision": 0}).status_code, 422)
        self.assertEqual(self.client.patch("/api/documents/missing", json={"fields": {}, "revision": 0}).status_code, 404)
        bad = self.client.post("/api/pipeline/run", files={"file": ("x.txt", b"test")}, data={"pipeline": "invalid"})
        self.assertEqual(bad.status_code, 400)
        self.assertEqual(self.upload(b"a" * (MAX_FILE_SIZE + 1)).status_code, 413)
        self.assertEqual(self.client.get("/api/documents").json()["documents"], [])

    def test_safe_download_and_range(self):
        response = self.upload(b"<script>alert(1)</script>", "../../danger.html", "text/html")
        document_id = response.json()["documentId"]
        source = self.client.get(f"/api/documents/{document_id}/original")
        self.assertIn("attachment", source.headers["content-disposition"])
        self.assertEqual(source.headers["x-content-type-options"], "nosniff")
        self.assertIn(("test-bucket", f"test/originals/{document_id}"), self.s3_client.objects)
        self.assertFalse((self.storage / "originals").exists())
        partial = self.client.get(f"/api/documents/{document_id}/original", headers={"Range": "bytes=0-6"})
        self.assertEqual(partial.status_code, 206)
        self.assertEqual(partial.content, b"<script")

    def test_s3_original_survives_restart_and_temporary_copy_is_deleted(self):
        document_id = self.upload(b"hello", "source.txt", "text/plain").json()["documentId"]
        with self.app.state.sessions() as session:
            self.assertEqual(session.get(Document, document_id).original_key, f"s3:test/originals/{document_id}")
        # A prefix change must not change the key of previously saved originals.
        self.originals.prefix = "new-prefix"
        with self.restarted_app() as client:
            with patch.object(self.originals, "download", wraps=self.originals.download) as download:
                response = client.get(f"/api/documents/{document_id}/original")
                self.assertEqual(response.content, b"hello")
                self.assertEqual(download.call_count, 1)
            with patch("backend.api.documents.TemporaryFileResponse", wraps=TemporaryFileResponse) as response_class:
                response = client.get(f"/api/documents/{document_id}/original", headers={"Range": "bytes=99-100"})
                self.assertEqual(response.status_code, 416)
                self.assertFalse(response_class.call_args.args[0].exists())

    def test_s3_failure_does_not_create_document_or_local_file(self):
        error = ClientError({"Error": {"Code": "AccessDenied", "Message": "secret"}}, "PutObject")
        with patch.object(self.s3_client, "put_object", side_effect=error):
            response = self.upload()
        self.assertEqual(response.status_code, 503)
        self.assertNotIn("secret", response.text)
        self.assertEqual(self.client.get("/api/documents").json()["documents"], [])
        self.assertFalse((self.storage / "originals").exists())

    def test_s3_object_is_removed_when_database_commit_fails(self):
        from sqlalchemy.exc import SQLAlchemyError
        with patch("sqlalchemy.orm.SessionTransaction.commit", side_effect=SQLAlchemyError("failed")):
            response = self.upload()
        self.assertEqual(response.status_code, 503)
        self.assertEqual(self.s3_client.objects, {})
        self.assertEqual(self.client.get("/api/documents").json()["documents"], [])

    def test_missing_and_unavailable_s3_originals(self):
        document_id = self.upload().json()["documentId"]
        url = f"/api/documents/{document_id}/original"
        self.s3_client.objects.clear()
        self.assertEqual(self.client.get(url).status_code, 404)
        error = ClientError({"Error": {"Code": "AccessDenied"}}, "GetObject")
        with patch.object(self.s3_client, "get_object", side_effect=error):
            self.assertEqual(self.client.get(url).status_code, 503)

    def test_legacy_local_original_remains_readable(self):
        document_id = self.upload(b"legacy", "old.txt", "text/plain").json()["documentId"]
        folder = self.storage / "originals"
        folder.mkdir()
        (folder / document_id).write_bytes(b"legacy")
        with self.app.state.sessions.begin() as session:
            session.get(Document, document_id).original_key = document_id
        response = self.client.get(f"/api/documents/{document_id}/original")
        self.assertEqual(response.content, b"legacy")
        self.assertTrue((folder / document_id).exists())

    def test_docx_text(self):
        from docx import Document
        doc = Document()
        doc.add_paragraph("Договор номер 15")
        buf = io.BytesIO()
        doc.save(buf)
        response = self.upload(buf.getvalue(), "contract.docx", "application/vnd.openxmlformats-officedocument.wordprocessingml.document")
        self.assertEqual(response.status_code, 200, response.text)
        self.assertIn("Договор номер 15", response.json()["text"])

    def submit_background(self, content=b"background text", pipeline=None):
        response = self.client.post("/api/pipeline/run?background=true", files={"file": ("test.txt", content, "text/plain")}, data={"pipeline": json.dumps(pipeline or {"source": "document"})})
        self.assertEqual(response.status_code, 202, response.text)
        with self.app.state.sessions() as session:
            event = session.scalar(select(TaskOutbox).where(TaskOutbox.job_id == response.json()["taskId"]))
            self.assertIsNotNone(event)
        return response.json()

    def execute_background(self, job_id):
        execute_job(job_id, self.app.state.sessions, self.originals, httpx.MockTransport(self.upstream), settings=JobSettings(max_attempts=1))

    def test_background_processing_and_redelivery_do_not_duplicate_documents(self):
        job = self.submit_background(pipeline={"source": "document", "extraction": {"mode": "fields", "model": "extract", "fields": [{"name": "total"}, {"name": "missing"}]}})
        self.assertEqual(self.calls, [])
        self.assertEqual((self.client.get(job["statusUrl"]).json()["status"], self.client.get(job["statusUrl"]).json()["stage"]), ("queued", "queued"))
        self.assertEqual(self.client.get("/api/documents").json()["documents"], [])
        self.execute_background(job["taskId"])
        self.execute_background(job["taskId"])
        status = self.client.get(job["statusUrl"]).json()
        self.assertEqual((status["status"], status["stage"], status["text"], status["documentId"]), ("succeeded", "completed", "background text", job["taskId"]))
        self.assertNotIn("failedStage", status)
        self.assertEqual(len(self.calls), 1)
        self.assertEqual(len(self.client.get("/api/documents").json()["documents"]), 1)
        document = self.client.get(f"/api/documents/{job['taskId']}").json()["document"]
        self.assertEqual(document["fields"], {"total": "1500", "missing": None})
        self.assertEqual(self.client.get(f"/api/documents/{job['taskId']}/original").content, b"background text")
        with self.restarted_app() as client:
            self.assertEqual(client.get(job["statusUrl"]).json()["status"], "succeeded")

    def test_background_failure_is_reported_and_can_be_retried(self):
        job = self.submit_background(pipeline={"source": "document", "extraction": {"mode": "prompt", "model": "extract"}})
        self.fail_upstream = True
        self.execute_background(job["taskId"])
        status = self.client.get(job["statusUrl"]).json()
        self.assertEqual(status["status"], "failed")
        self.assertEqual((status["stage"], status["failedStage"]), ("failed", "extraction"))
        self.assertIn("LiteLLM", status["error"])
        self.assertEqual(self.client.get("/api/documents").json()["documents"], [])
        self.fail_upstream = False
        self.assertEqual(self.client.post(job["statusUrl"] + "/retry").status_code, 202)
        self.execute_background(job["taskId"])
        retried = self.client.get(job["statusUrl"]).json()
        self.assertEqual((retried["status"], retried["stage"]), ("succeeded", "completed"))
        self.assertNotIn("failedStage", retried)
        self.assertEqual(self.client.post(job["statusUrl"] + "/retry").status_code, 409)

    def test_invalid_structured_output_is_not_saved_synchronously_or_by_worker(self):
        pipeline = {"source": "document", "extraction": {"mode": "fields", "model": "extract", "fields": [{"name": "total"}]}}
        invalid = '{"total":1500,"extra":"unrequested"}'
        with patch("backend.handler.litellm.complete", return_value=invalid):
            response = self.upload(pipeline=pipeline)
        self.assertEqual(response.status_code, 502)
        self.assertEqual(self.s3_client.objects, {})
        job = self.submit_background(pipeline=pipeline)
        with patch("backend.handler.litellm.complete", return_value=invalid):
            self.execute_background(job["taskId"])
        failed = self.client.get(job["statusUrl"]).json()
        self.assertEqual((failed["status"], failed["stage"], failed["failedStage"]), ("failed", "failed", "validation"))
        self.assertEqual(self.client.get("/api/documents").json()["documents"], [])

    def test_background_jobs_are_private_to_their_api_key(self):
        pipeline = self.create_pipeline("Счета")
        first = self.issue_key(**{"name": "ERP", "pipeline_ids": [pipeline]}).json()
        second = self.issue_key(**{"name": "CRM", "pipeline_ids": [pipeline]}).json()
        integration = TestClient(self.app, headers={"Authorization": "Bearer " + first["token"]})
        other = TestClient(self.app, headers={"Authorization": "Bearer " + second["token"]})
        response = self.send_document(integration, f"/api/v1/pipelines/{pipeline}/run?background=true")
        self.assertEqual(response.status_code, 202)
        job = response.json()
        self.assertTrue(job["statusUrl"].startswith("/api/v1/jobs/"))
        self.assertEqual(other.get(job["statusUrl"]).status_code, 404)
        self.assertEqual(other.post(job["statusUrl"] + "/retry").status_code, 404)
        self.assertEqual(TestClient(self.app).get(job["statusUrl"]).status_code, 401)
        replacement = self.create_pipeline("Другой")
        self.client.patch(f"/api/clients/{first['key']['client_id']}", json={"name": "ERP", "pipeline_ids": [replacement]})
        self.assertEqual(integration.post(job["statusUrl"] + "/retry").status_code, 404)
        self.execute_background(job["taskId"])
        self.assertEqual(integration.get(job["statusUrl"]).status_code, 404)
        own = integration.get("/api/v1/documents").json()["documents"]
        self.assertEqual(own, [])
        self.client.delete(f"/api/keys/{first['key']['id']}")
        self.assertEqual(integration.get(job["statusUrl"]).status_code, 401)

    def test_client_has_ten_active_background_job_slots_shared_by_all_keys(self):
        self.app.state.processing.settings = JobSettings(client_active_limit=10)
        pipeline = self.create_pipeline("Limited")
        issued = self.issue_key("Limited client", [pipeline]).json()
        client_id = issued["key"]["client_id"]
        second = self.client.post("/api/keys", json={
            "name": "Second key", "client_id": client_id,
            "permissions": ["run", "results", "history"],
        }).json()
        first_key = TestClient(self.app, headers={"Authorization": "Bearer " + issued["token"]})
        second_key = TestClient(self.app, headers={"Authorization": "Bearer " + second["token"]})

        failed = self.send_document(first_key, f"/api/v1/pipelines/{pipeline}/run?background=true").json()
        with self.app.state.sessions.begin() as session:
            session.get(ProcessingJob, failed["taskId"]).status = "failed"
        active = []
        for index in range(10):
            response = self.send_document(first_key if index < 5 else second_key, f"/api/v1/pipelines/{pipeline}/run?background=true")
            self.assertEqual(response.status_code, 202, response.text)
            active.append(response.json()["taskId"])

        stored_objects = len(self.s3_client.objects)
        blocked = self.send_document(second_key, f"/api/v1/pipelines/{pipeline}/run?background=true")
        self.assertEqual(blocked.status_code, 429, blocked.text)
        self.assertEqual(blocked.headers["Retry-After"], "5")
        self.assertIn("10 активных задач", blocked.json()["error"])
        self.assertEqual(len(self.s3_client.objects), stored_objects)
        self.assertEqual(second_key.post(f"/api/v1/jobs/{failed['taskId']}/retry").status_code, 429)

        # Another client has an independent pool of ten slots.
        other = self.issue_key("Other client", [pipeline]).json()["token"]
        other_response = self.send_document(TestClient(self.app, headers={"Authorization": "Bearer " + other}), f"/api/v1/pipelines/{pipeline}/run?background=true")
        self.assertEqual(other_response.status_code, 202, other_response.text)

        self.execute_background(active[0])
        retry = second_key.post(f"/api/v1/jobs/{failed['taskId']}/retry")
        self.assertEqual(retry.status_code, 202, retry.text)

    def test_publish_failure_preserves_task_id_and_original_for_recovery(self):
        job = self.submit_background(b"recover")
        publisher = lambda *_: (_ for _ in ()).throw(ConnectionError("redis down"))
        self.assertEqual(dispatch_once(self.app.state.sessions, publisher=publisher), 0)
        self.assertEqual(self.client.get(job["statusUrl"]).json()["status"], "queued")
        with self.app.state.sessions.begin() as session:
            event = session.scalar(select(TaskOutbox).where(TaskOutbox.job_id == job["taskId"]))
            self.assertIsNone(event.published_at)
            event.next_attempt_at = 0
        published = []
        self.assertEqual(dispatch_once(self.app.state.sessions, publisher=lambda *args: published.append(args)), 1)
        self.assertEqual(published[0][0], job["taskId"])
        self.execute_background(job["taskId"])
        self.assertEqual(self.client.get(job["statusUrl"]).json()["text"], "recover")

    def test_background_persistence_failure_cleans_s3_and_does_not_publish(self):
        from sqlalchemy.exc import SQLAlchemyError
        with patch("sqlalchemy.orm.SessionTransaction.commit", side_effect=SQLAlchemyError("failed")):
            response = self.client.post("/api/pipeline/run?background=true", files={"file": ("test.txt", b"test", "text/plain")}, data={"pipeline": '{"source":"document"}'})
        self.assertEqual(response.status_code, 503)
        self.assertEqual(self.s3_client.objects, {})
        with self.app.state.sessions() as session:
            self.assertEqual(list(session.scalars(select(TaskOutbox))), [])

    @unittest.skipUnless(os.getenv("TEST_DATABASE_URL"), "PostgreSQL required for concurrency test")
    def test_concurrent_worker_deliveries_are_serialized(self):
        from concurrent.futures import ThreadPoolExecutor
        job = self.submit_background(pipeline={"source": "document", "extraction": {"mode": "fields", "model": "extract", "fields": [{"name": "total"}]}})
        with ThreadPoolExecutor(max_workers=2) as pool:
            list(pool.map(self.execute_background, [job["taskId"], job["taskId"]]))
        self.assertEqual(len(self.calls), 1)
        self.assertEqual(len(self.client.get("/api/documents").json()["documents"]), 1)

    def test_redelivery_recovers_a_job_left_processing_by_a_lost_worker(self):
        job = self.submit_background()
        with self.app.state.sessions.begin() as session:
            session.get(ProcessingJob, job["taskId"]).status = "processing"
        self.execute_background(job["taskId"])
        self.assertEqual(self.client.get(job["statusUrl"]).json()["status"], "succeeded")

    @unittest.skipUnless(os.getenv("TEST_DATABASE_URL"), "PostgreSQL required for real worker test")
    def test_real_celery_worker_processes_a_queued_document(self):
        import time
        from celery import Celery
        from celery.contrib.testing.worker import start_worker
        from backend.worker.tasks import process_document
        test_celery = Celery("ocr_test", broker="memory://")
        test_celery.conf.update(task_ignore_result=True, task_serializer="json", accept_content=["json"])
        task = test_celery.task(process_document.run, name="ocr.process_document")
        job = self.submit_background()
        url = self.database_url.render_as_string(hide_password=False)
        with patch.dict("os.environ", {"DATABASE_URL": url}), start_worker(test_celery, pool="solo", perform_ping_check=False, loglevel="WARNING"):
            task.apply_async(args=[job["taskId"]])
            deadline = time.monotonic() + 10
            while time.monotonic() < deadline:
                status = self.client.get(job["statusUrl"]).json()
                if status["status"] in {"succeeded", "failed"}:
                    break
                time.sleep(0.05)
            self.assertEqual(status["status"], "succeeded", status)
            self.assertEqual(status["text"], "background text")
        test_celery.close()

    def test_outbox_survives_api_restart_without_publication(self):
        job = self.submit_background()
        with self.restarted_app() as client:
            sent = []
            self.assertEqual(dispatch_once(self.app.state.sessions, publisher=lambda *args: sent.append(args)), 1)
            self.assertEqual(sent[0][:2], (job["taskId"], 0))
            self.execute_background(job["taskId"])
            self.assertEqual(client.get(job["statusUrl"]).json()["status"], "succeeded")

    def test_published_events_of_finished_jobs_are_pruned(self):
        import time
        from backend.db.infra.repositories import OutboxRepository
        settings = JobSettings()
        repository = OutboxRepository(self.app.state.sessions, settings)
        queued = self.submit_background()
        finished = self.submit_background()
        with self.app.state.sessions.begin() as session:
            session.get(ProcessingJob, finished["taskId"]).status = "succeeded"
        moment = time.time()
        for _ in range(2):
            event = repository.claim(clock=moment)
            repository.finish(event, True, clock=moment)
        # Свежие события не трогаем: их ещё может потребоваться перевыпустить.
        self.assertEqual(repository.prune(clock=moment), 0)
        self.assertEqual(repository.prune(clock=moment + settings.lease + 1), 1)
        with self.app.state.sessions() as session:
            remaining = session.scalars(select(TaskOutbox.job_id)).all()
        self.assertEqual(remaining, [queued["taskId"]])

    def test_unknown_job_status_and_stage_are_rejected_by_the_database(self):
        from sqlalchemy.exc import IntegrityError
        job = self.submit_background()
        for column, value in (("status", "half-done"), ("stage", "thinking"), ("failed_stage", "thinking")):
            with self.subTest(column=column), self.assertRaises(IntegrityError):
                with self.app.state.sessions.begin() as session:
                    session.execute(update(ProcessingJob).where(ProcessingJob.id == job["taskId"]).values(**{column: value}))

    def test_outbox_recovers_lost_dispatcher_and_fences_stale_receipts(self):
        import time
        from backend.db.infra.repositories import OutboxRepository
        job = self.submit_background()
        repository = OutboxRepository(self.app.state.sessions, JobSettings())
        first = repository.claim()
        self.assertIsNone(repository.claim())
        second = repository.claim(clock=time.time() + 31)
        self.assertEqual(second.id, first.id)
        repository.finish(first, True)
        with self.app.state.sessions() as session:
            self.assertIsNone(session.get(TaskOutbox, first.id).published_at)
        repository.finish(second, True)

    def test_worker_lease_prevents_parallel_claim_and_stale_completion(self):
        import time
        from backend.db.infra.repositories import JobRepository
        from backend.schema.pipeline import Pipeline
        from backend.service.processing import build_document
        job = self.submit_background()
        repository = JobRepository(self.app.state.sessions, JobSettings())
        first = repository.claim(job["taskId"], clock=time.time() - 301)
        second = repository.claim(job["taskId"])
        self.assertIsNotNone(second)
        self.assertIsNone(repository.claim(job["taskId"]))
        def document(claimed):
            return build_document(claimed.id, claimed.original_key, claimed.filename, claimed.mime_type, claimed.size, Pipeline.model_validate(claimed.pipeline), "text", "text", claimed.pipeline_id, claimed.api_key_id)
        self.assertFalse(repository.complete(first, document(first)))
        self.assertTrue(repository.complete(second, document(second)))
        self.assertEqual(len(self.client.get("/api/documents").json()["documents"]), 1)

    def test_expired_worker_is_requeued_and_old_generation_is_ignored(self):
        import time
        from backend.db.infra.repositories import JobRepository, OutboxRepository
        job = self.submit_background()
        repository = JobRepository(self.app.state.sessions, JobSettings())
        old = repository.claim(job["taskId"], clock=time.time() - 301)
        OutboxRepository(self.app.state.sessions, JobSettings()).recover()
        self.assertIsNone(repository.claim(job["taskId"], generation=old.generation))
        with self.app.state.sessions() as session:
            row = session.get(ProcessingJob, job["taskId"])
            self.assertEqual((row.status, row.generation), ("queued", 1))
            self.assertEqual(len(list(session.scalars(select(TaskOutbox).where(TaskOutbox.job_id == row.id)))), 2)
        execute_job(job["taskId"], self.app.state.sessions, self.originals, httpx.MockTransport(self.upstream), generation=1)
        self.assertEqual(self.client.get(job["statusUrl"]).json()["status"], "succeeded")

    def test_temporary_model_failure_retries_with_backoff_then_succeeds(self):
        import time
        pipeline = {"source": "document", "extraction": {"mode": "prompt", "model": "extract"}}
        job = self.submit_background(pipeline=pipeline)
        self.fail_upstream = True
        execute_job(job["taskId"], self.app.state.sessions, self.originals, httpx.MockTransport(self.upstream))
        with self.app.state.sessions.begin() as session:
            row = session.get(ProcessingJob, job["taskId"])
            self.assertEqual((row.status, row.attempts, row.generation), ("queued", 1, 1))
            self.assertGreater(row.next_run_at, time.time())
            row.next_run_at = 0
        self.fail_upstream = False
        execute_job(job["taskId"], self.app.state.sessions, self.originals, httpx.MockTransport(self.upstream), generation=1)
        self.assertEqual(self.client.get(job["statusUrl"]).json()["status"], "succeeded")

    def test_attempt_budget_and_deadline_stop_automatic_retries(self):
        import time
        job = self.submit_background(pipeline={"source": "document", "extraction": {"mode": "prompt", "model": "extract"}})
        self.fail_upstream = True
        for _ in range(3):
            with self.app.state.sessions.begin() as session:
                session.get(ProcessingJob, job["taskId"]).next_run_at = 0
            execute_job(job["taskId"], self.app.state.sessions, self.originals, httpx.MockTransport(self.upstream))
        self.assertEqual(self.client.get(job["statusUrl"]).json()["status"], "failed")
        expired = self.submit_background()
        with self.app.state.sessions.begin() as session:
            row = session.get(ProcessingJob, expired["taskId"])
            row.status, row.deadline_at, row.lease_until = "processing", time.time() - 1, time.time() + 100
        from backend.db.infra.repositories import OutboxRepository
        OutboxRepository(self.app.state.sessions, JobSettings()).recover()
        self.assertEqual(self.client.get(expired["statusUrl"]).json()["status"], "failed")

    def test_model_io_occurs_without_database_transaction(self):
        from sqlalchemy import event
        engine = self.app.state.sessions.kw["bind"]
        active = [0]
        def begin(_):
            active[0] += 1
        def end(_):
            active[0] -= 1
        for name, callback in (("begin", begin), ("commit", end), ("rollback", end)):
            event.listen(engine, name, callback)
            self.addCleanup(event.remove, engine, name, callback)
        job = self.submit_background()
        async def process(*args, **kwargs):
            self.assertEqual(active[0], 0)
            return "text", "text"
        with patch("backend.service.processing.process_pipeline", side_effect=process):
            self.execute_background(job["taskId"])
        self.assertEqual(active[0], 0)
        self.assertEqual(self.client.get(job["statusUrl"]).json()["status"], "succeeded")

    def test_throttle_is_shared_with_a_second_api_instance_and_restart(self):
        browser = TestClient(self.app)
        for _ in range(MAX_FAILED_LOGINS):
            browser.post("/api/auth/login", json={"login": ADMIN_LOGIN, "password": "wrong-password"})
        with TestClient(self.make_app()) as other:
            response = other.post("/api/auth/login", json=ADMIN_CREDENTIALS)
            self.assertEqual(response.status_code, 429)

    def test_sync_timeout_returns_gateway_timeout_without_persistence(self):
        import asyncio
        self.app.state.processing.settings = JobSettings(timeout=1)
        async def slow(*args):
            await asyncio.sleep(5)
        with patch("backend.service.processing.recognize", side_effect=slow):
            response = self.upload()
        self.assertEqual(response.status_code, 504)
        self.assertEqual(self.s3_client.objects, {})

    def test_proxy_credentials_are_hidden_in_responses_and_logs(self):
        self.refuse_upstream = True
        self.app.state.proxy = "http://proxy-user:proxy-password@proxy.local:8080"
        with self.assertLogs("ocr", level="WARNING") as logs:
            response = self.upload(b"image", "scan.png", "image/png", {"source": "scans", "ocr": {"provider": "litellm", "model": "vision"}})
        combined = response.text + "\n".join(logs.output)
        self.assertIn("proxy.local:8080", combined)
        self.assertNotIn("proxy-user", combined)
        self.assertNotIn("proxy-password", combined)


    @unittest.skipUnless(os.getenv("TEST_DATABASE_URL"), "PostgreSQL required for JSONB migration")
    def test_jsonb_migration_preserves_existing_values_and_supports_rollback(self):
        from alembic import command
        from alembic.config import Config
        from sqlalchemy.dialects.postgresql import JSONB, JSON
        from backend.db.migrate import migrate
        columns = [("documents", "fields"), ("processing_jobs", "pipeline"), ("pipelines", "config"), ("integration_clients", "pipeline_ids"), ("login_guard", "failures")]
        payload = {"id":"jsonb_test", "name":"JSONB тест", "source":"document"}
        created = self.client.post("/api/pipelines", json=payload)
        self.assertEqual(created.status_code, 201, created.text)
        self.assertEqual(self.issue_key(**{"name":"JSONB key", "pipeline_ids":[created.json()["pipeline"]["id"]]}).status_code, 201)
        self.assertEqual(self.upload(b'{"nested":{"amount":1500,"items":[true,null,"text"]}}').status_code, 200)
        self.submit_background(pipeline={"name":"JSONB задача", "source":"document"})
        self.client.post("/api/auth/login", json={"login":"wrong", "password":"wrong"})
        engine = self.app.state.sessions.kw["bind"]

        def snapshot():
            with engine.connect() as connection:
                return {table:connection.execute(sql(f"SELECT {column} FROM {table} ORDER BY {column}::text")).scalars().all() for table, column in columns}

        def assert_types(expected):
            for table, column in columns:
                actual = next(item["type"] for item in inspect(engine).get_columns(table) if item["name"] == column)
                self.assertIsInstance(actual, expected)

        before = snapshot()
        assert_types(JSONB)
        config = Config()
        config.set_main_option("script_location", str(Path(__file__).resolve().parents[1] / "db" / "migration"))
        with engine.begin() as connection:
            config.attributes["connection"] = connection
            command.downgrade(config, "0002_reliable_jobs")
        assert_types(JSON)
        self.assertEqual(snapshot(), before)
        migrate(engine)
        assert_types(JSONB)
        self.assertEqual(snapshot(), before)

    def test_history_pipeline_filter_paginates_and_preserves_deleted_history(self):
        for pipeline_id in ("first", "second", "all"):
            self.assertEqual(self.client.post("/api/pipelines/import", json=[{"id":pipeline_id, "name":"Same name", "source":"document"}]).status_code, 200)
        first = self.client.post("/api/pipelines/first/run", files={"file":("first.txt", b"first", "text/plain")}).json()["documentId"]
        self.client.post("/api/pipelines/second/run", files={"file":("second.txt", b"second", "text/plain")})
        self.upload(b"Without pipeline", "test.txt", "text/plain")
        with self.app.state.sessions.begin() as session:
            original = session.get(Document, first)
            values = {column.name:getattr(original, column.name) for column in Document.__table__.columns if column.name != "id"}
            session.add_all([Document(id=str(uuid4()), **values) for _ in range(50)])
        page = self.client.get("/api/documents?pipeline_id=first").json()
        self.assertEqual(len(page["documents"]), 50)
        self.assertTrue(page["hasMore"])
        self.assertTrue(all(row["pipeline_id"] == "first" for row in page["documents"]))
        last = self.client.get("/api/documents?pipeline_id=first&page=1").json()
        self.assertEqual(len(last["documents"]), 1)
        self.assertFalse(last["hasMore"])
        self.assertFalse({row["id"] for row in page["documents"]} & {row["id"] for row in last["documents"]})
        self.assertEqual(len(self.client.get("/api/documents?pipeline_id=second").json()["documents"]), 1)
        self.assertEqual(self.client.get("/api/documents?pipeline_id=unknown").json()["documents"], [])
        self.assertEqual(self.client.get("/api/documents?pipeline_id=all").json()["documents"], [])
        self.client.delete("/api/pipelines/first")
        choices = self.client.get("/api/history/pipelines").json()["pipelines"]
        self.assertEqual({row["id"] for row in choices}, {"first", "second", "all"})
        self.assertTrue(next(row for row in choices if row["id"] == "first")["deleted"])
        self.assertEqual(len(self.client.get("/api/documents?pipeline_id=first").json()["documents"]), 50)

    def test_history_search_and_page_size(self):
        ids = []
        for filename in ("Invoice-100%.txt", "Invoice-200.txt", "Other.txt"):
            response = self.upload(b"Document", filename, "text/plain")
            self.assertEqual(response.status_code, 200, response.text)
            ids.append(response.json()["documentId"])
        first = self.client.get("/api/documents", params={"q": "invoice", "page_size": 1}).json()
        second = self.client.get("/api/documents", params={"q": "invoice", "page_size": 1, "page": 1}).json()
        self.assertTrue(first["hasMore"])
        self.assertFalse(second["hasMore"])
        self.assertEqual({first["documents"][0]["id"], second["documents"][0]["id"]}, set(ids[:2]))
        literal = self.client.get("/api/documents", params={"q": "%"}).json()
        self.assertEqual([row["id"] for row in literal["documents"]], [ids[0]])
        exact = self.client.get("/api/documents", params={"q": ids[2]}).json()
        self.assertEqual([row["id"] for row in exact["documents"]], [ids[2]])
        self.assertEqual(self.client.get("/api/documents", params={"q": "missing"}).json()["documents"], [])
        self.assertNotIn("text", first["documents"][0])
        self.assertEqual(self.client.get("/api/documents?page_size=101").status_code, 422)

    def test_history_pipeline_filter_keeps_integration_ownership(self):
        pipeline_id = self.create_pipeline("Owned")
        self.client.post(f"/api/pipelines/{pipeline_id}/run", files={"file":("admin.txt", b"Admin document", "text/plain")})
        token = self.issue_key(**{"name":"Integration", "pipeline_ids":[pipeline_id]}).json()["token"]
        self.client.cookies.clear()
        headers = {"Authorization":f"Bearer {token}"}
        uploaded = self.client.post(f"/api/v1/pipelines/{pipeline_id}/run", headers=headers, files={"file":("owned.txt", b"Own document", "text/plain")})
        self.assertEqual(uploaded.status_code, 200)
        rows = self.client.get(f"/api/v1/documents?pipeline_id={pipeline_id}", headers=headers).json()["documents"]
        self.assertEqual([row["id"] for row in rows], [uploaded.json()["documentId"]])
        self.assertEqual(self.client.get("/api/v1/documents?q=admin&page_size=20", headers=headers).json()["documents"], [])
        self.assertEqual(len(self.client.get("/api/v1/documents?q=owned&page_size=20", headers=headers).json()["documents"]), 1)
        self.assertEqual(self.client.get("/api/history/pipelines", headers=headers).status_code, 403)

    def test_background_status_links_match_api_prefix(self):
        for prefix in ("/api", "/api/v1"):
            response = self.client.post(f"{prefix}/pipeline/run?background=true", files={"file":("test.txt", b"Test", "text/plain")}, data={"pipeline":'{"source":"document"}'})
            self.assertEqual(response.status_code, 202, response.text)
            job = response.json()
            self.assertEqual(job["statusUrl"], f"{prefix}/jobs/{job['taskId']}")
            self.assertEqual(response.headers["Location"], job["statusUrl"])
            self.assertEqual(self.client.get(job["statusUrl"]).status_code, 200)

    def test_migrated_schema_matches_models(self):
        from alembic.autogenerate import compare_metadata
        from alembic.migration import MigrationContext
        from backend.db.domain.models import Base
        with self.app.state.sessions.kw["bind"].connect() as connection:
            context = MigrationContext.configure(connection, opts={"compare_type": True})
            self.assertEqual(compare_metadata(context, Base.metadata), [])

    def test_combined_processing_settings_survive_edit_and_run(self):
        config = {"id":"combined", "name":"Combined", "source":"document", "extraction":{"model":"extract", "prompt_enabled":True, "fields_enabled":True, "prompt":"Use the final total", "fields":[{"name":"total"}]}}
        response = self.client.post("/api/pipelines", json=config)
        self.assertEqual(response.status_code, 201, response.text)
        saved = response.json()["pipeline"]
        result = self.upload(b"Total 1500", "test.txt", "text/plain", saved)
        self.assertEqual(result.status_code, 200, result.text)
        self.assertEqual(json.loads(result.json()["result"]), {"total":"1500"})
        self.assertIn("Use the final total", json.loads(self.calls[-1].content)["messages"][0]["content"])
        saved["extraction"].update(prompt_enabled=False, fields_enabled=False)
        changed = self.client.patch(f"/api/pipelines/{saved['id']}", json=saved)
        self.assertEqual(changed.status_code, 200, changed.text)
        with self.restarted_app() as restarted:
            stored = restarted.get("/api/pipelines").json()["pipelines"][0]
            self.assertFalse(stored["extraction"]["prompt_enabled"])
            self.assertFalse(stored["extraction"]["fields_enabled"])
            self.assertEqual(stored["extraction"]["prompt"], "Use the final total")
            self.assertEqual(stored["extraction"]["fields"], [{"name":"total", "description":"", "public_description":"", "type":"string", "fields":[]}])
        before = len(self.calls)
        job = self.submit_background(pipeline=stored)
        self.execute_background(job["taskId"])
        self.assertEqual(self.client.get(job["statusUrl"]).json()["status"], "succeeded")
        self.assertEqual(len(self.calls), before)

    def test_disabled_ocr_reads_text_without_calling_vision(self):
        config = {"source":"scans", "ocr":{"provider":"litellm", "enabled":False, "model":"vision", "prompt":"Saved OCR prompt"}}
        result = self.upload(b"Original text", "test.txt", "text/plain", config)
        self.assertEqual(result.status_code, 200, result.text)
        self.assertEqual(result.json()["result"], "Original text")
        self.assertEqual(self.calls, [])

    @unittest.skipUnless(os.getenv("TEST_DATABASE_URL"), "PostgreSQL required for shared row locking")
    def test_concurrent_login_guards_do_not_lose_failures(self):
        from concurrent.futures import ThreadPoolExecutor
        from backend.service.throttle import LoginThrottle
        guards = [LoginThrottle(self.app.state.sessions), LoginThrottle(self.app.state.sessions)]
        with ThreadPoolExecutor(max_workers=4) as pool:
            results = list(pool.map(lambda index: guards[index % 2].verify(lambda: False), range(16)))
        self.assertEqual(sum(retry == 0 for accepted, retry in results), MAX_FAILED_LOGINS)
        self.assertEqual(sum(retry > 0 for accepted, retry in results), 6)
        with self.app.state.sessions() as session:
            self.assertEqual(len(session.get(LoginGuard, "admin").failures), MAX_FAILED_LOGINS)

    @unittest.skipUnless(os.getenv("TEST_DATABASE_URL"), "PostgreSQL required for concurrent migrations")
    def test_parallel_migration_processes_adopt_a_fresh_schema(self):
        import subprocess
        import sys
        target = self.database_for(self.storage)
        environment = {**os.environ, "DATABASE_URL": target.render_as_string(hide_password=False)}
        processes = [subprocess.Popen([sys.executable, "-m", "backend.db.migrate"], env=environment, stdout=subprocess.PIPE, stderr=subprocess.PIPE) for _ in range(2)]
        try:
            for process in processes:
                stdout, stderr = process.communicate(timeout=30)
                self.assertEqual(process.returncode, 0, stderr.decode(errors="replace"))
            engine, _ = open_database(target)
            try:
                with engine.connect() as connection:
                    self.assertEqual(connection.execute(sql("SELECT version_num FROM alembic_version")).scalar_one(), "0005_keycloak")
                    self.assertEqual(connection.execute(sql("SELECT count(*) FROM login_guard")).scalar_one(), 1)
            finally:
                engine.dispose()
        finally:
            for process in processes:
                if process.poll() is None:
                    process.kill()
                    process.communicate()


    def test_public_nested_contract_sync_background_and_history(self):
        config = {"id": "typed", "name": "Typed", "description": "Public pipeline", "source": "document", "extraction": {
            "model": "extract", "prompt": "SECRET PROMPT", "fields": [{"name": "objects", "type": "array",
            "description": "SECRET RULE", "public_description": "Objects", "fields": [{"name": "count", "type": "integer", "public_description": "Count"}]}]}}
        created = self.client.post("/api/pipelines", json=config)
        self.assertEqual(created.status_code, 201, created.text)
        config["id"] = created.json()["pipeline"]["id"]
        token = self.issue_key(**{"name": "Typed", "pipeline_ids": [config["id"]]}).json()["token"]
        integration = TestClient(self.app, headers={"Authorization": "Bearer " + token})
        contract = integration.get("/api/v1/pipelines").json()["pipelines"][0]
        self.assertEqual(set(contract), {"id", "name", "description", "executionModes", "asyncConcurrency", "resultSchema"})
        self.assertNotIn("SECRET", json.dumps(contract))
        self.assertEqual(contract["resultSchema"]["properties"]["objects"]["items"]["properties"]["count"]["type"], ["integer", "null"])
        self.assertEqual(integration.get("/api/pipelines").json()["pipelines"], [contract])
        self.assertIn("extraction", self.client.get("/api/pipelines").json()["pipelines"][0])
        expected = {"objects": [{"count": 3}]}
        from unittest.mock import AsyncMock
        with patch("backend.handler.litellm.complete", new=AsyncMock(return_value=json.dumps(expected))):
            response = self.send_document(integration, f"/api/v1/pipelines/{config['id']}/run")
            self.assertEqual(response.status_code, 200, response.text)
            self.assertEqual(response.json()["result"], expected)
            doc_id = response.json()["documentId"]
            queued = self.send_document(integration, f"/api/v1/pipelines/{config['id']}/run?background=true").json()
            self.app.state.processing.execute(queued["taskId"])
        self.assertEqual(integration.get(f"/api/v1/jobs/{queued['taskId']}").json()["result"], expected)
        self.assertEqual(integration.get(f"/api/v1/documents/{doc_id}").json()["document"]["result"], expected)
        # Result contract survives config changes and edits to the separate operator fields.
        self.client.patch(f"/api/pipelines/{config['id']}", json={**config, "extraction": None})
        self.client.patch(f"/api/documents/{doc_id}", json={"fields": {}, "revision": 0})
        self.assertEqual(integration.get(f"/api/v1/documents/{doc_id}").json()["document"]["result"], expected)
        self.assertEqual(integration.get("/api/v1/pipelines").json()["pipelines"][0]["resultSchema"], {"type": "string"})
        raw = integration.post(f"/api/v1/pipelines/{config['id']}/run", files={"file": ("x.txt", b'{"hello":123}', "text/plain")})
        self.assertEqual(raw.json()["result"], '{"hello":123}')
        self.assertEqual(integration.get(f"/api/v1/documents/{raw.json()['documentId']}").json()["document"]["result"], '{"hello":123}')


    def test_client_ownership_rotation_and_current_pipeline_access(self):
        pipeline = self.create_pipeline("Client pipeline")
        issued = self.issue_key("Owner", [pipeline]).json()
        client_id = issued["key"]["client_id"]
        def credential(token):
            return TestClient(self.app, headers={"Authorization": "Bearer " + token})
        original = credential(issued["token"])
        rotated = self.client.post("/api/keys", json={"name": "Replacement", "client_id": client_id, "permissions": ["run", "results", "history"]}).json()
        replacement = credential(rotated["token"])
        foreign = credential(self.issue_key("Other", [pipeline]).json()["token"])
        doc_id = self.send_document(original, f"/api/v1/pipelines/{pipeline}/run").json()["documentId"]
        job = self.send_document(original, f"/api/pipelines/{pipeline}/run?background=true").json()
        self.client.delete(f"/api/keys/{issued['key']['id']}")
        self.assertEqual(original.get(f"/api/v1/documents/{doc_id}").status_code, 401)
        self.assertEqual(replacement.get(f"/api/v1/documents/{doc_id}").status_code, 200)
        self.assertEqual(foreign.get(f"/api/v1/documents/{doc_id}").status_code, 404)
        self.assertEqual(foreign.get(job["statusUrl"]).status_code, 404)
        self.assertEqual(foreign.post(job["statusUrl"] + "/retry").status_code, 404)
        self.assertEqual(replacement.post(job["statusUrl"] + "/retry").status_code, 202)
        self.execute_background(job["taskId"])
        self.assertEqual(replacement.get(job["statusUrl"]).json()["status"], "succeeded")
        with self.app.state.sessions() as session:
            self.assertEqual(session.get(Document, job["taskId"]).client_id, client_id)
            self.assertEqual(session.get(Document, doc_id).client_id, client_id)
            self.assertEqual(session.get(ProcessingJob, job["taskId"]).client_id, client_id)
        self.assertEqual(len(replacement.get("/api/v1/documents").json()["documents"]), 2)
        self.client.patch(f"/api/clients/{client_id}", json={"name": "Owner", "pipeline_ids": []})
        self.assertEqual(replacement.get("/api/v1/documents").json()["documents"], [])
        self.assertEqual(replacement.get(f"/api/v1/documents/{doc_id}").status_code, 404)
        for prefix in ("/api", "/api/v1"):
            self.assertEqual(replacement.get(f"{prefix}/jobs/{job['taskId']}").status_code, 404)
            self.assertEqual(replacement.post(f"{prefix}/jobs/{job['taskId']}/retry").status_code, 404)
            self.assertEqual(self.send_document(replacement, f"{prefix}/pipelines/{pipeline}/run").status_code, 403)
        self.assertEqual(self.client.get(f"/api/documents/{doc_id}").status_code, 200)
        self.client.patch(f"/api/clients/{client_id}", json={"name": "Owner", "pipeline_ids": [pipeline]})
        self.assertEqual(replacement.get(f"/api/v1/documents/{doc_id}").status_code, 200)
        self.client.delete(f"/api/pipelines/{pipeline}")
        self.assertEqual(replacement.get(f"/api/v1/documents/{doc_id}").status_code, 404)
        self.assertEqual(replacement.get(job["statusUrl"]).status_code, 404)

    def test_key_action_permissions_cover_all_response_paths(self):
        pipeline = self.create_pipeline("Permission pipeline")
        issued = self.issue_key("Permissions", [pipeline]).json()
        identity = {"name": "Permissions", "client_id": issued["key"]["client_id"]}
        integration = TestClient(self.app, headers={"Authorization": "Bearer " + issued["token"]})
        job = self.send_document(integration, f"/api/v1/pipelines/{pipeline}/run?background=true").json()
        self.execute_background(job["taskId"])
        doc_url = f"/api/v1/documents/{job['taskId']}"
        for permission in ("run", "results", "history"):
            with self.subTest(permission=permission):
                changed = self.client.patch(f"/api/keys/{issued['key']['id']}", json={**identity, "permissions": [permission]})
                self.assertEqual(changed.status_code, 200, changed.text)
                self.assertEqual(integration.get(doc_url).status_code, 200 if permission == "results" else 403)
                self.assertEqual(integration.get("/api/v1/documents").status_code, 200 if permission == "history" else 403)
                for prefix in ("/api", "/api/v1"):
                    self.assertEqual(integration.get(f"{prefix}/jobs/{job['taskId']}").status_code, 200 if permission == "results" else 403)
                    for url, data in ((f"{prefix}/pipelines/{pipeline}/run", None), (f"{prefix}/pipeline/run", {"pipeline_id": pipeline})):
                        response = self.send_document(integration, url, data)
                        self.assertEqual(response.status_code, 200 if permission == "run" else 403)
                        if permission == "run":
                            self.assertEqual(set(response.json()), {"file", "documentId", "status"})
                    self.assertEqual(integration.post(f"{prefix}/jobs/{job['taskId']}/retry").status_code, 409 if permission == "run" else 403)
        self.client.patch(f"/api/keys/{issued['key']['id']}", json={**identity, "permissions": []})
        self.assertEqual(integration.get(doc_url).status_code, 403)
        self.assertEqual(integration.get("/api/v1/documents").status_code, 403)

    def test_key_client_is_required_immutable_and_admin_managed(self):
        pipeline = self.create_pipeline("Access")
        issued = self.issue_key("First", [pipeline]).json()
        other = self.issue_key("Second", [pipeline]).json()
        identity = {"name": "Key", "client_id": issued["key"]["client_id"], "permissions": ["run"]}
        for payload in ({"name": "No client", "permissions": ["run"]}, {**identity, "client_id": "missing"}, {**identity, "permissions": ["admin"]}, {**identity, "pipeline_ids": [pipeline]}):
            self.assertEqual(self.client.post("/api/keys", json=payload).status_code, 422)
        self.assertEqual(self.client.patch(f"/api/keys/{issued['key']['id']}", json={**identity, "client_id": other["key"]["client_id"]}).status_code, 409)
        integration = TestClient(self.app, headers={"Authorization": "Bearer " + issued["token"]})
        for api_client, expected in ((TestClient(self.app), 401), (integration, 403)):
            self.assertEqual(api_client.get("/api/clients").status_code, expected)
            self.assertEqual(api_client.post("/api/clients", json={"name": "New", "pipeline_ids": []}).status_code, expected)
            self.assertEqual(api_client.patch(f"/api/clients/{identity['client_id']}", json={"name": "Changed", "pipeline_ids": []}).status_code, expected)
        self.assertEqual(len(self.client.get("/api/clients").json()["clients"]), 2)


class DatabaseConfigurationTests(unittest.TestCase):
    def test_runtime_rejects_sqlite_and_requires_postgres_credentials(self):
        with patch.dict("os.environ", {"DATABASE_URL": "sqlite:///data/ocr.sqlite3"}):
            with self.assertRaisesRegex(RuntimeError, "PostgreSQL"):
                configured_database_url()
        with patch.dict("os.environ", {"DATABASE_URL": "", "POSTGRES_PASSWORD": ""}):
            with self.assertRaisesRegex(RuntimeError, "POSTGRES_PASSWORD"):
                configured_database_url()

    def test_password_special_characters_are_preserved(self):
        with patch.dict("os.environ", {"DATABASE_URL": "", "POSTGRES_PASSWORD": "a@:/?#%$", "POSTGRES_HOST": "db"}):
            url = configured_database_url()
        self.assertEqual(url.password, "a@:/?#%$")
        self.assertEqual(url.host, "db")
        self.assertEqual(url.drivername, "postgresql+psycopg")

    def test_postgres_url_uses_psycopg3(self):
        with patch.dict("os.environ", {"DATABASE_URL": "postgres://ocr:test@127.0.0.1/ocr"}):
            engine, _ = open_database(configured_database_url())
        try:
            self.assertEqual(engine.dialect.driver, "psycopg")
        finally:
            engine.dispose()


class ProxyTests(unittest.TestCase):
    def setUp(self):
        blank = dict.fromkeys(("PROXY_URL", "HTTP_PROXY", "HTTPS_PROXY", "http_proxy", "https_proxy", "NO_PROXY"), "")
        self.environment = patch.dict("os.environ", blank)
        self.environment.start()
        self.addCleanup(self.environment.stop)

    def test_bare_address_gets_a_scheme_and_wins_over_the_environment(self):
        os.environ["HTTP_PROXY"] = "old.local:3128"
        os.environ["PROXY_URL"] = "proxy.company.local:8080"
        self.assertEqual(proxy_options(), "http://proxy.company.local:8080")
        self.assertEqual(os.environ["HTTP_PROXY"], "http://proxy.company.local:8080")
        self.assertEqual(os.environ["HTTPS_PROXY"], "http://proxy.company.local:8080")

    def test_environment_proxy_stays_the_fallback(self):
        os.environ["HTTPS_PROXY"] = "old.local:3128"
        self.assertIsNone(proxy_options())
        self.assertEqual(os.environ["HTTPS_PROXY"], "http://old.local:3128")

    def test_unsupported_scheme_is_ignored(self):
        os.environ["PROXY_URL"] = "ftp://proxy.company.local:21"
        with self.assertLogs("ocr", level="WARNING"):
            self.assertIsNone(proxy_options())
        self.assertEqual(os.environ["HTTP_PROXY"], "")

    def test_ocr_service_requests_are_routed_through_the_proxy(self):
        os.environ["PROXY_URL"] = "proxy.company.local:8080"
        os.environ["NO_PROXY"] = "127.0.0.1"
        proxy_options()
        with httpx.Client() as client:
            routed = client._transport_for_url(httpx.URL("https://ocr.example.com/recognize"))
            excluded = client._transport_for_url(httpx.URL("http://127.0.0.1:8000/recognize"))
        self.assertEqual((routed._pool._proxy_url.host, routed._pool._proxy_url.port), (b"proxy.company.local", 8080))
        self.assertFalse(hasattr(excluded._pool, "_proxy_url"))




if __name__ == "__main__":
    unittest.main()
