import base64
import io
import json
import logging
import os
from typing import Literal

import httpx
from fastapi import HTTPException
from pydantic import BaseModel, Field, HttpUrl, model_validator
from .security import safe_url

logger = logging.getLogger("ocr")


class ExtractionField(BaseModel):
    name: str
    description: str = ""


class OCR(BaseModel):
    enabled: bool = True
    provider: Literal["litellm", "service"]
    temperature: float = Field(default=0, ge=0, le=2)
    model: str | None = None
    prompt: str | None = None
    url: HttpUrl | None = None
    # Поля формы конкретного OCR-сервиса (например model_name и force_ocr) — уходят вместе с файлом.
    options: dict[str, str] = Field(default_factory=dict)


class Extraction(BaseModel):
    mode: Literal["prompt", "fields"] = "fields"
    model: str = ""
    prompt_enabled: bool | None = None
    fields_enabled: bool | None = None
    temperature: float = Field(default=0, ge=0, le=2)
    max_tokens: int = Field(default=2048, ge=1, le=128000)
    prompt: str = ""
    fields: list[ExtractionField] = Field(default_factory=list)

    @model_validator(mode="after")
    def valid_field_names(self):
        if self.fields_enabled is not None:
            self.mode = "fields" if self.fields_enabled else "prompt"
        if self.active and not self.model.strip():
            raise ValueError("Выберите модель извлечения")
        if self.prompt_enabled is True and not self.prompt.strip():
            raise ValueError("Укажите промпт")
        if self.use_fields:
            names = [field.name for field in self.fields]
            if not names or any(not name.strip() for name in names) or len(names) != len(set(names)):
                raise ValueError("Укажите непустые уникальные имена полей извлечения")
        return self

    @property
    def use_fields(self):
        return self.fields_enabled if self.fields_enabled is not None else self.mode == "fields"

    @property
    def use_prompt(self):
        return self.prompt_enabled if self.prompt_enabled is not None else self.mode == "prompt"

    @property
    def active(self):
        return self.use_fields or self.use_prompt


class Pipeline(BaseModel):
    name: str = "Untitled"
    source: Literal["scans", "document"]
    ocr: OCR | None = None
    extraction: Extraction | None = None


SERVICE_TEXT_KEYS = ("page_content", "text", "result")
PROXY_VARIABLES = ("HTTP_PROXY", "HTTPS_PROXY", "http_proxy", "https_proxy")
PROXY_SCHEMES = ("http://", "https://", "socks5://", "socks5h://")


def normalized_proxy(value):
    # Corporate Windows environments sometimes supply proxy addresses without a scheme.
    value = (value or "").strip()
    if not value:
        return None
    return value if "://" in value else f"http://{value}"


def proxy_options():
    """PROXY_URL is the single knob for the corporate proxy; HTTP(S)_PROXY keep working as a fallback.

    The value is written back into the standard proxy variables, so httpx routes every outgoing
    request (external OCR service and LiteLLM alike) through it while still honouring NO_PROXY.
    """
    for name in PROXY_VARIABLES:
        value = normalized_proxy(os.getenv(name))
        if value:
            os.environ[name] = value
    proxy = normalized_proxy(os.getenv("PROXY_URL"))
    if proxy and not proxy.startswith(PROXY_SCHEMES):
        logger.warning("PROXY_URL=%s пропущен: поддерживаются схемы http, https, socks5.", safe_url(proxy))
        return None
    if proxy:
        for name in PROXY_VARIABLES:
            os.environ[name] = proxy
    return proxy


def litellm_config():
    base = os.getenv("LITELLM_BASE_URL", "").rstrip("/")
    if not base:
        raise HTTPException(503, "LITELLM_BASE_URL не настроен в .env")
    if not base.startswith(("http://", "https://")):
        raise HTTPException(503, "LITELLM_BASE_URL должен начинаться с http:// или https://")
    key = os.getenv("LITELLM_API_KEY")
    return base, {"Authorization": f"Bearer {key}"} if key else {}


def api_url(base, path):
    return f"{base}{'' if base.endswith('/v1') else '/v1'}/{path}"


def upstream_failure(message, status):
    error = HTTPException(502, message)
    error.transient = status in {408, 429, 500, 502, 503, 504}
    return error


async def complete(client: httpx.AsyncClient, body: dict):
    base, headers = litellm_config()
    response = await client.post(api_url(base, "chat/completions"), headers=headers, json=body)
    structured = body.get("response_format", {}).get("type") == "json_schema"
    if not response.is_success:
        if structured and response.status_code in {400, 422}:
            raise HTTPException(502, f"LiteLLM отклонил Structured Output (HTTP {response.status_code}). Проверьте поддержку json_schema выбранной моделью и настройки LiteLLM.")
        raise upstream_failure(f"LiteLLM вернул HTTP {response.status_code}. Проверьте модель и доступ к сервису.", response.status_code)
    try:
        choice = response.json()["choices"][0]
        message = choice["message"]
        if structured:
            if message.get("refusal"):
                raise HTTPException(502, "Модель отказалась извлекать поля документа")
            if choice.get("finish_reason") != "stop":
                raise HTTPException(502, "Модель не завершила Structured Output. Проверьте лимит max_tokens и ответ модели.")
        result = message["content"]
        if not isinstance(result, str):
            raise ValueError("Not text")
        return result
    except (ValueError, KeyError, IndexError, TypeError, AttributeError) as error:
        raise HTTPException(502, "LiteLLM вернул некорректный ответ") from error


def digital_text(content: bytes, filename: str, mime: str):
    try:
        if mime == "application/pdf" or filename.lower().endswith(".pdf"):
            from pypdf import PdfReader
            text = "\n\n".join(page.extract_text() or "" for page in PdfReader(io.BytesIO(content)).pages)
            if not text.strip():
                raise HTTPException(422, "В PDF нет текстового слоя. Выберите источник «Сканы / изображения».")
            return text.strip()
        if filename.lower().endswith(".docx"):
            from docx import Document
            doc = Document(io.BytesIO(content))
            rows = [paragraph.text for paragraph in doc.paragraphs]
            rows.extend("\t".join(cell.text for cell in row.cells) for table in doc.tables for row in table.rows)
            return "\n".join(rows).strip()
        return content.decode("utf-8-sig").strip()
    except HTTPException:
        raise
    except Exception as error:
        raise HTTPException(422, "Не удалось прочитать документ. Для изображений выберите источник «Сканы / изображения», для текста используйте UTF-8.") from error


async def recognize(client, pipeline, content, filename, mime):
    if pipeline.source == "document":
        from starlette.concurrency import run_in_threadpool
        return await run_in_threadpool(digital_text, content, filename, mime)
    ocr = pipeline.ocr
    if ocr is None:
        raise HTTPException(422, "Выберите способ распознавания сканов")
    if not ocr.enabled:
        from starlette.concurrency import run_in_threadpool
        return await run_in_threadpool(digital_text, content, filename, mime)
    if ocr.provider == "service":
        if not ocr.url:
            raise HTTPException(422, "Укажите URL OCR-сервиса")
        response = await client.post(str(ocr.url), files={"file": (filename, content, mime)}, data=ocr.options or None)
        if not response.is_success:
            raise upstream_failure(f"OCR-сервис вернул HTTP {response.status_code}", response.status_code)
        try:
            payload = response.json()
        except ValueError:
            return response.text
        text = next((payload[key] for key in SERVICE_TEXT_KEYS if key in payload), response.text) if isinstance(payload, dict) else response.text
        if not isinstance(text, str):
            raise HTTPException(502, "OCR-сервис должен вернуть текст в поле page_content, text или result")
        return text
    if not ocr.model:
        raise HTTPException(422, "Выберите vision-модель")
    return await complete(client, {
        "model": ocr.model, "temperature": ocr.temperature, "max_tokens": 4096,
        "messages": [{"role": "user", "content": [
            {"type": "text", "text": ocr.prompt or "Распознай весь текст на изображении."},
            {"type": "image_url", "image_url": {"url": f"data:{mime};base64,{base64.b64encode(content).decode('ascii')}"}},
        ]}],
    })


async def extract(client, extraction, text):
    if not extraction.active:
        return text
    structured = extraction.use_fields
    schema = {
        "type": "object",
        "properties": {field.name: {"type": ["string", "null"], "description": field.description} for field in extraction.fields},
        "required": [field.name for field in extraction.fields],
        "additionalProperties": False,
    }
    instruction = extraction.prompt if extraction.use_prompt else ""
    if structured:
        instruction += "\nИзвлеки значения полей документа согласно JSON Schema. Верни все заданные поля; значения — строки. Если значение не найдено, верни null. Не придумывай отсутствующие данные."
    result = await complete(client, {
        "model": extraction.model, "temperature": extraction.temperature, "max_tokens": extraction.max_tokens,
        **({"response_format": {"type": "json_schema", "json_schema": {"name": "document_fields", "strict": True, "schema": schema}}} if structured else {}),
        "messages": [
            {"role": "system", "content": f"{instruction}\n" + ("Отвечай только валидным JSON без пояснений." if structured else "Верни ответ единым текстом. Не создавай JSON-объект и именованные поля результата.")},
            {"role": "user", "content": text},
        ],
    })
    if structured:
        validate_extracted_fields(result, schema)
    return result


def validate_extracted_fields(result, schema):
    # The generated schema is flat and permits only string/null values. Validate
    # it locally too: a gateway may ignore the provider's response_format.
    def unique_object(pairs):
        obj = {}
        for key, value in pairs:
            if key in obj:
                raise ValueError("Duplicate JSON key")
            obj[key] = value
        return obj

    try:
        value = json.loads(result, object_pairs_hook=unique_object)
        if not isinstance(value, dict) or set(value) != set(schema["required"]):
            raise ValueError("Incorrect field names")
        if any(item is not None and not isinstance(item, str) for item in value.values()):
            raise ValueError("Incorrect field type")
    except (ValueError, TypeError) as error:
        raise HTTPException(502, "Ответ модели не соответствует JSON Schema: нужны все заданные поля, без лишних, со значениями строка или null.") from error


def result_fields(result):
    try:
        value = json.loads(result)
        return value if isinstance(value, dict) else {"result": value}
    except ValueError:
        return {}
